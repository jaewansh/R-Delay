#!/usr/bin/env python3
"""Create a code-aligned delivery without overwriting v0.1 or changing the engine.
Document version 0.3 is not a claim that the entire product is finished.
Run from the repository root after pytest and browser checks in the delivery workflow.
"""
from __future__ import annotations
import copy, hashlib, json, os, platform, re, subprocess, sys, zipfile
from pathlib import Path
from datetime import datetime, timezone
from xml.etree import ElementTree as ET
import fitz
from docx import Document
from docx.shared import Mm, Pt
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
import build_documents as g

ROOT=Path(__file__).resolve().parents[1]
BASE='4791514f711a0eab07a7b64c2773b85658737674'
PREP='1d1d61f68d4fdeb3e2fa48b9b34965e4b013f623'
BRANCH='codex/rdelay-final-delivery'
OUT=ROOT/'deliverables/final_20261022'
ART=ROOT/'artifacts/final_delivery'
for p in (OUT,ART,ART/'rendered',ART/'build_tmp'): p.mkdir(parents=True,exist_ok=True)
g.OUT=OUT;g.ART=ART;g.RENDER=ART/'rendered';g.TMP=ART/'build_tmp'
STAMP=datetime.now(timezone.utc).isoformat()
SOURCE=subprocess.check_output(['git','rev-parse','HEAD'],text=True).strip()

def git(*args): return subprocess.check_output(['git',*args],cwd=ROOT)
def save_json(p,x): p.write_text(json.dumps(x,ensure_ascii=False,indent=2),encoding='utf-8')
def digest(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def paragraphs(d):
    for p in d._element.xpath('.//w:p'): yield Paragraph(p,d)
    for sec in d.sections:
        yield from sec.header.paragraphs
        yield from sec.footer.paragraphs

def replace_text(d,mapping):
    for p in paragraphs(d):
        if p._p.xpath('.//m:oMath'): continue
        old=p.text;new=old
        for a,b in mapping.items(): new=new.replace(a,b)
        if new==old: continue
        runs=p.runs
        if runs:
            runs[0].text=new
            for r in runs[1:]: r.text=''
        else:p.add_run(new)

def restamp(d,label):
    d.core_properties.title=label
    d.core_properties.subject='2026-10-22 통합 전달본 / 실제 구현 v0.2와 이번 대회 목표 구분'
    for sec in d.sections:
        sec.header.paragraphs[0].text='R-DELAY / '+label+' / 문서 v0.3'
        for r in sec.header.paragraphs[0].runs:r.font.size=Pt(7)
        for r in sec.footer.paragraphs[0].runs:
            r.text=r.text.replace('1차 제출·설계 문서','통합 전달본 · 개인정보 직접 기입')

# Read fresh test evidence, never infer success from test source presence.
def execution_evidence():
    tree=ET.parse(ART/'pytest.xml').getroot()
    cases=tree.findall('.//testcase')
    summary={'tests':len(cases),'failures':len(tree.findall('.//failure')),
             'errors':len(tree.findall('.//error')),'skipped':len(tree.findall('.//skipped'))}
    if not cases or summary['failures'] or summary['errors'] or summary['skipped']:
        raise RuntimeError('The complete stable-code test suite must pass before delivery')
    browser=json.loads((ROOT/'artifacts/runtime/browser_check.json').read_text())
    if not browser.get('passed'):raise RuntimeError('Actual browser check did not pass')
    save_json(ART/'browser_check.json',browser)
    return {'generated_at':STAMP,'delivery_source_commit':SOURCE,'stable_code_commit':BASE,
            'pytest':summary,'browser':browser,'field_validation':'not performed',
            'public_deployment':'not performed','pyodide_execution':'not implemented',
            'human_visual_review':'not performed in this chat runtime'}


def final_application(evidence):
    g.SUMMARY='SIS 미분방정식으로 해상·육상 지연을 연결하고 로그-OU·입자필터·선석 대기를 통합한 웹 시제품이다. 관측 기반 회복률을 사용해 원래 마감 아래 배차 위험·비용·실행조건을 비교한다. 현재는 합성자료로 검산했으며 실제 항만 성능은 별도 검증한다.'
    d=g.application()
    mapping={
      '이번 대회 목표는 수리모델·계수 추정·웹/API의 통합 서비스, 같은 Python 엔진의 로컬 실행환경, 검증 기록과 GitHub 제출 버전이다. SIS·로그-OU·입자필터·선석 대기·배차 비교를 모두 이번 개발 범위로 둔다. 현재 문서는 1차 설계이며 이 기능들의 구현·실행 검증이 완료된 상태는 아니다.':
      '2거점 합성 웹·Python API 시제품과 SIS·로그-OU·입자필터·선석 대기·배차 계산을 구현했다. 회복률 하나의 추정과 자동검사를 포함한다. 본선에서는 외생변수 학습, 관측 확장, 사용성·민감도 검증을 보강한다. GitHub Pages의 브라우저 직접 계산은 전환 목표이며 아직 미구현이다.',
      '공개 근거·관측/단위·입출력 명세, 데이터 이용권·사전개발 허용 범위 확인':'합성 웹 시제품·코드·검산 확보. 실제 자료 이용권과 사전개발 허용 범위 재확인',
      'SIS·OU·필터·대기·계수 추정 통합, API에서 동일 결과 반환':'통합 엔진의 조건·관측 정의 검토, 외생변수 계수 학습과 시간순 평가 보강',
      '웹에 관측 전후 상태·위험·후보 사유 표시, 온라인/로컬 실행 점검':'웹 사용성·시나리오 보강, 브라우저 직접계산 호환성·로컬 시연 점검',
      '해석해·시간간격·확률보존·제약 검사, 시간순 평가·민감도 및 발표 정리':'상태·확률·제약·기준모형 비교, 시드·입자수 민감도 및 발표자료 확정',
      '현재 이용 가능한 설명변수에서 계수를 추정하고 이후 구간 평가. 상수 추정과 학습 모델 구분.':'합성 관측에서 상수 회복률 1개 추정 구현. 기상·물동량 기반 학습은 이번 대회 추가 구현 목표.',
      '지금 단계의 모델 수치시험·웹 배포는 미실행이다.':
      f'현재 합성 시제품의 수리모델·API {evidence["pytest"]["tests"]}개 검사와 실제 브라우저 검사를 수행했다. 공개 배포·현장 성능은 미검증이다.',
      '미래 관측·사후 평활화는 실시간 경보에 넣지 않는다.':'시제품은 10:00 단일 관측만 지원하고 다른 시각은 거부한다. 미래 정보는 넣지 않는다.',
      '□ SIS 미분방정식을 이번 구현의 핵심 엔진으로 유지 [U]':'□ SIS 미분방정식을 구현 엔진과 제출 본문에 유지 [U, C1]',
      '[U] 이번 사용자 제공 설계 기준.':'[U] 사용자 원안·설계 기준. [C1] jaewansh/R-Delay, 코드 4791514f.',
    }
    replace_text(d,mapping)
    restamp(d,'참가신청서·컨셉기획안')
    return d


def final_report(evidence, example):
    # Preserve the detailed, previously authored mathematics; update status and
    # implementation-specific contracts instead of relabeling an old draft.
    d=g.report()
    mapping={
      '종합설계서':'종합 기획·수리모델 설명서',
      'v0.1 / 1차 제출·설계 문서':'문서 v0.3 / 현재 구현을 반영한 통합 전달본',
      '공식 신청서와 일치하는 수식·입출력·검증·사업화 명세':'신청서·실제 코드 v0.2와 일치하는 수식·입출력·검증·사업화 설명',
      '실제 Word/PDF, 문서 생성 원본, 문서 환경·페이지 검사 기록':'갱신 Word/PDF, 전체 구현 코드 ZIP, 실제 재실행 검사·전달 기록',
      '이번에 구현하지 않는 것':'현재 미완료 범위',
      '서비스·수리모델 계산 엔진 전체. 이들은 이번 대회 구현 목표로 유지':'외생변수 머신러닝, 일반 관측 스트리밍, GitHub Pages/Pyodide 전환, 공개 배포·현장 검증',
      '작성 기준 2026-10-21. 수치 실적이나 모델 테스트 통과 수를 기재하지 않는다.':'작성 기준 2026-10-22. 현재 소스와 이번 재실행 검사를 기준으로 상태를 갱신했다. 계획과 합성 구현 검사를 현장 성과와 구분한다.',
      '18–20':'18–20 및 부록21–25',
      '사업모델·제출 점검·출처와 문서 검사':'사업모델·제출 점검·출처, 실제 코드·파라미터·재실행 결과·전달 한계',
      '이 문서 단계':'현재 구현 상태',
      '사용자 식을 보존한 설계':'원안 식을 보존한 2거점 코드·검사',
      '구현 목표 / 미구현':'2거점 합성 조건에서 구현·검사',
      '구현 목표 / 미실행':'상수 회복률 추정 구현; 외생변수 학습 미구현',
      '구현 목표 / 미배포':'로컬 웹/API 구현·검사; 공개 배포 미완료',
      '원문을 저장소에서 읽었다고 주장하지 않음. 개인실적을 인용하지 않음.':'이번 대화에 제공된 원안 텍스트·페이지를 확인. 개인 연구실적은 인용하지 않음.',
      '사용자 핵심식은 이번 메시지[U]를 기준으로 삼는다.':'핵심식은 사용자 원안·설계 기준[U]을, 구현값은 동봉 소스[C1]를 기준으로 삼는다.',
      '08  로그-OU 확률변동 — 구체화할 설계':'08  로그-OU 확률변동 — 구현과 가정',
      '위 식은 사용자 로그-OU 목표를 구체화한 추가 구현 명세다.':'위 식은 사용자 로그-OU 목표를 구체화한 구현식이다. engine/filtering.py에서 정확한 OU 끝점 전이를 사용한다.',
      '이번 문서에서는 미실행된 복원 수치를 검증 실적으로 싣지 않는다.':'복원값은 이번 전달 과정에서 다시 계산한 값만 부록에 싣는다. 지정 오차 범위는 신뢰구간이 아니다.',
      '입자 r의 SIS·로그-OU 상태와 연결된 대기열 상태/분포':'입자 r의 SIS·로그-OU 상태. 10:00에서 알려진 대기열 상태로 초기화',
      '재표본추출 시 상태뿐 아니라 같은 입자에 조건부인 대기열 분포/시뮬레이션 이력도 함께 복사한다. 다만 아직 태깅하지 않은 미래 대상 선박을 배경 시스템에 중복 삽입하지 않는다. 관측 갱신 요약은 재표본추출 전 가중 분포 기준으로 저장한다.':
      '현재 코드는 10:00에서 관측 갱신 후 대기열 선박 수가 직접 알려졌다고 초기화한다. 그 이전 대기열을 필터링하지 않는다. 이후 대기열은 입자별 미래 SIS 경로를 따른다. 관측 후 요약은 재표본추출 전 가중 분포이며 대상 선박을 한 번만 삽입한다.',
      '늦게 도착한 과거 관측은 원시시각부터 재처리하거나 재실행 필요 상태로 기록한다.':'현재는 10:00 관측 하나만 지원한다. 늦은/다른 시각 관측은 거부하며 일반 재처리는 미구현이다.',
      '실제 관측을 읽고 추정한 실행 결과는 이번에 제시하지 않는다.':'합성 관측에 대한 실행·복원 결과는 부록에 제시한다. 실제 항만 관측으로 추정한 것은 아니다.',
      '이번 학습 구현 목표는':'아직 미구현인 이번 대회 학습 목표는',
      '관측/고정조건 → 추정 모델 ID·진단·학습 구간':'합성 관측/고정조건 → 상수 회복률·진단 (학습 모델 저장소 없음)',
      '시나리오/모델/마감/후보 → 실행 ID·처리 상태':'유효한 합성 입력 → 동기 계산 결과·실행 ID (작업 큐 미구현)',
      '새 관측 → 원본을 보존하는 새 실행 ID':'동일 10:00 관측 교체 → 같은 시드 재계산·새 실행 ID',
      '긴 계산은 실행 ID로 진행상태를 조회하고 중복 요청을 구분한다.':'현재 계산은 동기 처리하며 동시 계산 1개로 제한한다. 비동기 작업 큐는 미구현이다.',
      '늦은 과거 관측은 재처리 규칙과 시간을 보존한다. 위 API는 이번 개발 명세이며 현재 배포된 주소가 아니다.':'다른 시각 관측은 현재 거부한다. API는 로컬 서버에서 구현·검사됐지만 공개 URL은 없다.',
      'React/TypeScript·FastAPI·컨테이너 구성을 개발 후보로 삼되 지금 설치·배포했다고 표현하지 않는다.':'실제 화면은 HTML/CSS/JavaScript, API는 FastAPI다. React는 사용하지 않는다. 컨테이너 설정은 있으나 Docker 실행 검사는 미수행이다.',
      '현재 이 모델 검사들은 미실행이며 문서 검사만 별도로 수행한다.':'구현된 검사와 실제 실행 결과는 부록21–24에 구분했다. 장기 스트리밍·현장 검증·일반 민감도 전 범위가 모두 완료된 것은 아니다.',
      '문서 생성과 검사는 실행환경 기록에서 확인한다.':'문서·코드 재실행 검사와 남은 육안 검수는 동봉 검사 기록에서 확인한다.',
      '이번 작업에서는 실제 모델 수치시험·웹 배포를 수행하지 않는다.':'현재 합성 모델·로컬 웹/API 검사는 수행했다. 공개 배포·현장 검증은 수행하지 않았다.',
    }
    replace_text(d,mapping)
    # Explicitly clear any obsolete blanket status statement not caught above.
    for p in paragraphs(d):
        if p._p.xpath('.//m:oMath'):continue
        if '모델 수치시험' in p.text and '미실행' in p.text:
            p.text='모델·API 재실행 검사와 브라우저 검사 결과는 부록 및 동봉 기록을 따른다. 공개 배포·현장 검증과 문서 전 페이지 직접 육안 검수는 미완료다.'
    g.page(d,'21  이번 전달본의 실제 구현·시험 상태')
    g.para(d,'문서 v0.3은 현재 구현된 서비스 v0.2를 설명하는 통합 기획서다. 제품 기능 전체나 현장 검증을 모두 끝냈다는 의미의 최종본은 아니다. 코드는 아래 고정 버전을 변경하지 않고 패키징했다.')
    g.para(d,'기준 코드 커밋 [C1]: '+BASE,9)
    g.table(d,[['구분','확인된 상태'],['SIS·로그-OU·입자필터','2거점, 10:00 단일 관측 합성 시나리오 구현'],['대기·준비·육상·배차','시간가변 선석, 대상 선박, 사건 순서·마감·제약 연결'],['계수 추정','과거 합성 관측에서 육상 회복률 1개 추정'],['웹/API','HTML/CSS/JS + FastAPI. Python 단일 엔진'],['이번 재실행 pytest',f'{evidence["pytest"]["tests"]}개 / 실패0·오류0·건너뜀0'],['브라우저','실제 로컬 API와 Chromium 검사. 상세 항목은 browser_check.json'],['미완료','외생변수 학습, 임의시각 스트리밍, Pyodide, 공개 배포·현장 검증']],[44,132],True,9.3)
    g.para(d,'검사 통과는 합성 계산과 연결의 일관성을 확인한다. 실제 납기 예측 정확도·비용 절감·고객 구매 의사를 입증하지 않는다. 코드에 없는 기능을 계획에서 삭제하지 않되 구현된 기능으로 기재하지 않는다.')

    g.page(d,'22  코드에 실제 사용된 파라미터와 초기분포')
    g.table(d,[['항목','현재 기본값·단위','성격'],['시간','08:00 기준 t=0; 판단2h; 선박도착3h; 원래마감8h; 종료16h','가상 2026-10-21 사례B'],['적분 간격','1/60 h (1분)','수치설정'],['작업집단 N','해상100, 육상100','합성 일정 작업 단위'],['초기 지연 비율','해상 Beta(20,20), 육상 Beta(3,47)','입자별 합성 사전분포'],['기준 전파율','해상0.8/h, 육상0.6/h','고정 합성값'],['회복률','해상0.4/h; 육상 관측으로 추정','추정 조건부'],['결합','C_LP=0.5κ; 대각1; 역방향0','κ 기본1, 범위0~3'],['OU','θ=0.6/h; z 평균0; σ기본0.25/h^(1/2)','z(0) 정규 정상분산 σ²/(2θ)'],['입자','기본256; 64/128/256/512 지원','ESS<P/2 재표본추출'],['난수','시드20261021; 5개 독립 스트림 분기','재현 설정']],[34,93,49],True,8.7)
    g.para(d,'해석해 검사에서의 해상50%·육상0% 초기조건과 실제 웹 시나리오의 Beta 초기분포는 다르다. 기존 13:15 마감 예제의 위험값을 이 16:00 마감 시나리오에 강제로 재사용하지 않는다.',9.5)
    g.para(d,'확률변동은 계수 변동의 합성 모델이다. 실제 태풍 발생 시각·강도를 예측한 결과가 아니다. 지수변환 β의 기준값은 평균 β와 자동으로 같지 않다.',9.5)

    g.page(d,'23  대기·관측·비용 파라미터와 적용 제한')
    g.table(d,[['항목','현재 값·단위 / 의미'],['현재 지연 관측','기본 육상 20개 중8개, 집중도φ=25. 10:00 베타-이항 갱신'],['신규 지연 사건','08~10시 신규 생성항 적분, 음이항 크기12, 보고율1'],['알려진 선박 수','10:00에 시스템 내2척. 대기열 관측동화는 아님'],['선석','동질 서버2개; 배경도착0.4척/h; 기준서비스1.5/h/선석'],['SIS 서비스 영향','η_B=0.8; μ=1.5 exp(-0.8 x_P) × 개방상태'],['유한 대기 상태','기본0~40 및 흡수 넘침상태; 범위8~80'],['화물 준비','선박 작업완료+0.25h와 문서준비11:15 중 늦은 시각'],['육상 처리','h0=0.5h, η_L=0.8; 완료는 정규화 작업량1 도달'],['도로 이동','고정1.5h. 시간대별 교통학습 없음'],['비용','준비대기30,000원/h, 마감실패600,000원/건'],['후보','11:30유지0원; 12:30변경10,000원; 13:30변경15,000원'],['승인·접근','승인20분+여유10분+접근30분; 변경마감 별도'],['위험기준','기본20%; 입력 변경은 민감도 실험']],[44,132],True,8.5)
    g.para(d,'모든 값은 합성 가정 또는 계산 설정이다. 실측 견적·실제 고객 계약·공식 항만 처리율이 아니다. 준비 미완료 표본이나 대기열 넘침이 있으면 비용을 확정하지 않고 추천을 보류한다.',9.3)

    g.page(d,'24  이번 재계산의 결과와 해석')
    g.para(d,'다음 수치는 동봉 코드에 기본 입력을 실제로 전달해 생성했다. 과거 대화의 숫자를 복사한 표가 아니다. 하나의 시드에서 얻은 조건부 몬테카를로 표본값이며 현장 신뢰구간이 아니다.')
    g.table(d,[['항목','실제 계산값'],['육상 회복률',f'{example["calibration"]["gamma_land_per_h"]:.9f} /h'],['관측 전 육상 지연',f'{100*example["filter"]["prior_mean"][1]:.4f}%'],['관측 후 육상 지연',f'{100*example["filter"]["posterior_mean"][1]:.4f}%'],['유효입자수',f'{example["filter"]["ess"]:.3f} / {example["input"]["particles"]}'],['판정',example['decision']['label']]],[60,116],True,9.5)
    g.table(d,[['후보','표본 실패위험','표본 비교비용','판정'],*[[a['label'],f'{100*a["risk_upper"]:.3f}%',('미완결' if a['cost_krw'] is None else f'{a["cost_krw"]:,.0f}원'),('조건충족' if a['eligible'] else '제외/기준초과')] for a in example['candidates']]],[38,42,48,48],True,9)
    g.para(d,'risk_upper/lower는 구현상 상태범위 밖 경로를 처리하는 표본 경계다. 추정·파라미터·모델 오류까지 포함한 통계적 신뢰상하한이 아니다. 현재는 입자당 미래 대상선박 표본1개이므로 표본 수·시드 민감도 검증을 확대해야 한다.')
    g.para(d,'원본 결과: artifacts/final_delivery/example_result.json. 테스트: pytest.xml, pytest.log, browser_check.json. 모델 소스의 해시는 code_manifest.json에 기록한다. API의 임시 실행ID와 계산시간은 매 실행 달라질 수 있다.',9.3)

    g.page(d,'25  코드 사용법·배포 상태·최종 전달 조건')
    g.para(d,'전체 코드는 RDelay_전체코드_v0.2_20261022.zip의 R-Delay 폴더에 있다. 웹·API·엔진·검사·생성 스크립트·기록과 이전 문서를 보존한 저장소 스냅샷이다. 본 전달본의 기획서/PDF는 별도 문서 v0.3을 사용한다.')
    g.para(d,'실행: Python 3.12 환경에서 아래 명령을 순서대로 실행한다.',10,True)
    for line in ['python -m venv .venv','source .venv/bin/activate','python -m pip install -r requirements-service.txt','python -m uvicorn api.main:app --host 127.0.0.1 --port 8000']:
        g.para(d,line,9)
    g.para(d,'브라우저: http://127.0.0.1:8000. Windows PowerShell 활성화는 .venv\\Scripts\\Activate.ps1. HTML 더블클릭만으로 Python API가 실행되지는 않는다.',9.5)
    g.para(d,'GitHub Pages + Pyodide 전환은 합의한 검토 방향이며 아직 구현·검사되지 않았다. Render는 필수가 아니고, 별도 배포 준비 코드도 실제 공개 서버를 생성한 상태는 아니다. 이번 전달 작업에서는 서버를 만들거나 결제·배포하지 않는다.')
    g.para(d,'별도 ZIP RDelay_배포준비코드_미배포.zip은 PR #2 브랜치의 전체 소스 스냅샷이다. 기본 구현 코드와 혼합하지 않았으며, 실제 서비스 URL은 없다. GitHub Pages 호환 엔진으로 잘못 표시하지 않는다.')
    g.para(d,'전체 문서는 PDF로 변환하고 모든 페이지를 이미지로 생성하며 수식·개인정보 공란·쪽수·경계를 자동검사한다. 이 대화에서 전 페이지 이미지를 직접 보는 육안검수와 Microsoft Word 자체 검수는 미완료로 남긴다. 학적·서명·증빙과 최신 공지를 확인한 뒤 실제 제출한다.')
    g.para(d,'[C1] https://github.com/jaewansh/R-Delay/tree/'+BASE,8.5)
    g.para(d,'[C2] 배포 준비: https://github.com/jaewansh/R-Delay/tree/'+PREP,8.5)
    restamp(d,'종합 기획·수리모델 설명서')
    return d


def package_code(target,sha,label,evidence):
    raw=git('archive','--format=zip','--prefix=R-Delay/',sha)
    target.write_bytes(raw)
    # Do not redistribute local fonts or secrets. Existing repository snapshot
    # is checked, not augmented with credentials or user identification.
    with zipfile.ZipFile(target,'a',zipfile.ZIP_DEFLATED) as z:
        bad=[n for n in z.namelist() if Path(n).suffix.lower() in {'.ttf','.otf','.woff','.woff2'}]
        if bad:raise RuntimeError('Font distribution prohibited: '+repr(bad))
        z.writestr('DELIVERY_README_KO.md',f'''# {label}\n\n기준 커밋: {sha}\n문서 버전 v0.3과 서비스 코드 버전0.2는 별개다.\n\nR-Delay/에 실제 저장소 전체 소스가 있다. 기존 v0.1 문서는 이력을 보존한 것이며 최신 기획서는 별도 전달된 v0.3 파일이다.\n\n실행: R-Delay 폴더에서 `python -m pip install -r requirements-service.txt` 다음 `python -m uvicorn api.main:app --host 127.0.0.1 --port 8000`. 브라우저에서 http://127.0.0.1:8000.\n\nGitHub Pages/Pyodide 실행은 아직 구현되지 않았다. 공개 URL도 생성되지 않았다. 합성 시제품이며 현장 예측력·비용절감·고객 구매는 미검증이다.\n''')
        z.writestr('DELIVERY_TEST_SUMMARY.json',json.dumps(evidence,ensure_ascii=False,indent=2))
        if sha==BASE:
            for name in ['pytest.xml','pytest.log','browser_check.json','example_result.json']:
                p=ART/name
                if p.exists():z.write(p,'delivery_checks/'+name)
    with zipfile.ZipFile(target) as z:
        if z.testzip():raise RuntimeError('Corrupt ZIP')


def main():
    if git('branch','--show-current').decode().strip()!=BRANCH:raise RuntimeError('Wrong branch')
    evidence=execution_evidence()
    for name in ['engine','api','web','tests','requirements-service.txt']:
        if git('diff',BASE,'HEAD','--',name).strip():raise RuntimeError('Stable service source changed: '+name)
    from api.schemas import RunRequest
    from engine.simulation import run_simulation
    example=run_simulation(RunRequest().model_dump(mode='json'))
    save_json(ART/'example_result.json',example)
    save_json(ART/'execution_summary.json',evidence)
    app=final_application(evidence);report=final_report(evidence,example)
    names=['RDelay_최종기획서_참가신청서_v0.3.docx','RDelay_종합기획서_수리모델설명_v0.3.docx']
    app.save(OUT/names[0]);report.save(OUT/names[1])
    apdf=g.convert(OUT/names[0],OUT);rpdf=g.convert(OUT/names[1],OUT)
    ainfo=g.inspect_pdf(apdf,'application_final');rinfo=g.inspect_pdf(rpdf,'comprehensive_final')
    # Word edits must preserve editable equations and all personal cells blank.
    errors=[]
    for ri in range(1,5):
        for ci in (1,2,3,4,6):
            if app.tables[1].cell(ri,ci).text.strip():errors.append(f'{ri}:{ci}')
        if app.tables[1].cell(ri,5).text!='□재학 □휴학':errors.append('academic checkbox')
    native=len(app._element.xpath('.//m:oMath'))
    with fitz.open(apdf) as f:
        p3=f[2].get_text();math_visible=all(s in p3 for s in ['β','γ','∑'])
        app_text='\n'.join(p.get_text() for p in f)
    with fitz.open(rpdf) as f:report_text='\n'.join(p.get_text() for p in f)
    checks={'generated_at':STAMP,'source_commit':SOURCE,'stable_code_commit':BASE,
            'summary_characters':len(g.SUMMARY),'motivation_characters':len(g.MOTIVE),
            'keywords':len(g.KEYS),'editable_sis_equations':native,'privacy_errors':errors,
            'sis_pdf_symbols_present':math_visible,'application_pages':ainfo['pages'],
            'report_pages':rinfo['pages'],'application':ainfo,'report':rinfo,
            'visual_review':'not performed; all page renders generated',
            'word_native_render_review':'not performed; LibreOffice used',
            'local_skills':'Both /home/oai/skills docx and pdfs reads failed before container initialization; remote workflow cannot access them',
            'public_deployment':'not performed','pages_pyodide':'not implemented'}
    checks['automatic_pass']=(ainfo['pages']==5 and len(g.SUMMARY)<=200 and len(g.MOTIVE)<=200 and native>=3 and not errors and math_visible and all(not p['outside_page_spans'] for v in (ainfo,rinfo) for p in v['page_checks']))
    (ART/'application_text.txt').write_text(app_text,encoding='utf-8')
    (ART/'report_text.txt').write_text(report_text,encoding='utf-8')
    save_json(ART/'document_check.json',checks)
    (ART/'document_check.md').write_text('# 최종 전달 문서 검사\n\n'+json.dumps({k:v for k,v in checks.items() if k not in ('application','report')},ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    if not checks['automatic_pass']:raise RuntimeError('Document automatic checks failed: '+json.dumps({k:v for k,v in checks.items() if k not in ('application','report')},ensure_ascii=False))
    package_code(OUT/'RDelay_전체코드_v0.2_20261022.zip',BASE,'기본 실행 코드 전체',evidence)
    package_code(OUT/'RDelay_배포준비코드_미배포.zip',PREP,'별도 공개 배포 준비 소스 - 미배포',{'record':'기존 PR #2 사전검사 기록을 따름; 이번 기본코드40개 검사와 혼동 금지','source_commit':PREP,'public_deployment':False})
    manifest={'generated_at':STAMP,'document_version':'0.3','service_version':'0.2.0',
              'delivery_source_commit':SOURCE,'stable_code_commit':BASE,'deployment_preparation_commit':PREP,
              'files':[{'path':str(p.relative_to(ROOT)),'bytes':p.stat().st_size,'sha256':digest(p)} for p in sorted(OUT.iterdir()) if p.is_file()]}
    save_json(OUT/'manifest.json',manifest)
    (OUT/'먼저읽기_전달범위.md').write_text('# R-Delay 통합 전달본\n\n최종기획서 Word와 종합 PDF는 현재 구현 v0.2를 반영한 문서 v0.3이다. 서비스 전체 완성·공개배포·현장검증을 의미하지 않는다. 개인정보는 직접 입력한다.\n\n- 기본 코드: RDelay_전체코드_v0.2_20261022.zip\n- 배포준비 소스: 별도 ZIP, 미배포·main 미통합. Pages/Pyodide는 미구현.\n- 문서 전체 렌더링·자동검사 완료. 전 페이지 직접 육안검수·MS Word 확인 미완료.\n- 제출 시 서명·동의/서약·재학/휴학 증명은 별도다.\n- 기존 문서·main·서비스 소스 변경 없음.\n',encoding='utf-8')
    with zipfile.ZipFile(OUT/'RDelay_기획서와코드_전체전달.zip','w',zipfile.ZIP_DEFLATED) as z:
        for p in OUT.iterdir():
            if p.is_file() and p.name!='RDelay_기획서와코드_전체전달.zip':z.write(p,p.name)
        for p in [Path(__file__),ROOT/'scripts/build_documents.py',ART/'document_check.json',ART/'execution_summary.json']:
            z.write(p,str(p.relative_to(ROOT)))
    print(json.dumps(manifest,ensure_ascii=False,indent=2));print('DOCUMENT_PAGES',ainfo['pages'],rinfo['pages'])

if __name__=='__main__':main()
