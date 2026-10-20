#!/usr/bin/env python3
"""Verify real equation rendering, then build documents with a tested OMML form.
This wrapper is the actual workflow entry point. build_documents.py is its content
and layout module. No service/model test or human visual review is claimed.
"""
from __future__ import annotations
import copy
import json
import re
from pathlib import Path
import fitz
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import build_documents as b


def native_probe():
    md=b.TMP/'native_probe.md'; original=b.TMP/'native_probe.docx'
    md.write_text('$$\n'+b.EQUATIONS['SIS']+'\n$$\n',encoding='utf-8')
    b.run(['pandoc',str(md),'-o',str(original)])
    source=Document(original)
    math_paras=[p for p in source.paragraphs if p._p.xpath('.//m:oMath')]
    if not math_paras: raise RuntimeError('Pandoc failed to create a native equation paragraph.')
    tests=[]
    for variant in ('original','whole_paragraph','inline_math','display_math'):
        target=b.TMP/('probe_'+variant+'.docx')
        if variant=='original':
            source.save(target)
        else:
            doc=b.init_doc('수식 렌더 검사',True)
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            if variant=='whole_paragraph':
                p._p.addprevious(copy.deepcopy(math_paras[0]._p))
                p._p.getparent().remove(p._p)
            elif variant=='inline_math':
                for element in math_paras[0]._p.xpath('.//m:oMath'):
                    p._p.append(copy.deepcopy(element))
            else:
                for element in math_paras[0]._p.xpath('.//m:oMathPara'):
                    p._p.append(copy.deepcopy(element))
            doc.save(target)
        pdf=b.convert(target,b.TMP)
        with fitz.open(pdf) as f:
            text='\n'.join(p.get_text() for p in f)
            # Header/footer do not contain these mathematical symbols. Seeing
            # a native formula object in DOCX is not sufficient to pass.
            greek=sum(c in text for c in ('β','γ'))
            math_seen=greek==2 and ('∑' in text or '∑' in text.replace(' ',''))
            result={'variant':variant,'text':text,'math_seen':math_seen,'pages':len(f),'greek_symbol_count':greek}
            f[0].get_pixmap(matrix=fitz.Matrix(2,2),alpha=False).save(b.RENDER/('probe_'+variant+'.png'))
            tests.append(result)
    (b.ART/'math_render_probe.json').write_text(json.dumps(tests,ensure_ascii=False,indent=2),encoding='utf-8')
    working=[r['variant'] for r in tests if r['math_seen']]
    if not working:
        raise RuntimeError('All native math rendering probes failed; deliverables will not be rebuilt without equations. Ensure libreoffice-math is installed. Probe text: '+repr(tests))
    return 'whole_paragraph' if 'whole_paragraph' in working else ('inline_math' if 'inline_math' in working else 'display_math')


CACHE={}
def install_formula_renderer(variant):
    def equation(doc,key,number=None):
        if key not in CACHE:
            md=b.TMP/(key+'_native.md');target=b.TMP/(key+'_native.docx')
            md.write_text('$$\n'+b.EQUATIONS[key]+'\n$$\n',encoding='utf-8')
            b.run(['pandoc',str(md),'-o',str(target)])
            source=Document(target)
            found=[p for p in source.paragraphs if p._p.xpath('.//m:oMath')]
            if not found:raise RuntimeError('Missing native equation: '+key)
            CACHE[key]=copy.deepcopy(found[0]._p)
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(5)
        if variant=='whole_paragraph':
            source_p=copy.deepcopy(CACHE[key])
            # Preserve Pandoc's valid paragraph structure and all OMML.
            # Number is in its own following compact right-aligned paragraph.
            p._p.addprevious(source_p)
            if number:
                p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(3)
                p.add_run('('+str(number)+')').font.size=Pt(7)
            else:
                p._p.getparent().remove(p._p)
        else:
            nodes=CACHE[key].xpath('.//m:oMath') if variant=='inline_math' else CACHE[key].xpath('.//m:oMathPara')
            for node in nodes:p._p.append(copy.deepcopy(node))
            if number:p.add_run('  ('+str(number)+')').font.size=Pt(7)
        return p
    b.equation=equation


def main():
    variant=native_probe()
    install_formula_renderer(variant)
    b.main()
    app=b.OUT/'RDelay_참가신청서_v0.1.pdf'
    with fitz.open(app) as f:
        p3=f[2].get_text()
        seen=all(c in p3 for c in ('β','γ','∑'))
        if not seen:raise RuntimeError('Generated application page 3 did not retain visible extracted equation glyphs.')
    jpath=b.ART/'document_check.json'
    checks=json.loads(jpath.read_text())
    checks['native_render_variant']=variant
    checks['application_pdf_equation_glyphs_present']=seen
    checks['automatic_pass']=checks['automatic_pass'] and seen
    jpath.write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
    mpath=b.ART/'document_check.md'
    mpath.write_text(mpath.read_text()+'\n## 수식 실표시 추가 검사\n- LibreOffice Math 구성요소를 설치한 변환환경에서 OMML 원본·삽입 방식 4종을 검사했다.\n- 채택한 수식 구조: '+variant+'\n- 신청서 PDF 3쪽의 β·γ·합 기호 텍스트 존재: '+str(seen)+'\n- 이는 자동 추출 검사이며 육안검수 완료를 뜻하지 않는다.\n',encoding='utf-8')
    print('Verified formula rendering mode:',variant)

if __name__=='__main__':main()
