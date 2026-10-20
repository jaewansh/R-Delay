#!/usr/bin/env python3
"""Document-only R-Delay bootstrap. No model execution or deployment is claimed.
Run on the isolated document branch: python scripts/build_documents.py
Dependencies are installed by .github/workflows/documents.yml.
Original application/notice binaries were absent in the repository; use the audited
conversation copies and the user's explicit design brief, not an invented repository read.
"""
from __future__ import annotations
import copy
import hashlib
import json
import os
import platform
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from io import BytesIO

import fitz
from PIL import Image, ImageOps, ImageDraw, ImageFont
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'deliverables'; ART=ROOT/'artifacts'; DOCS=ROOT/'docs'
RENDER=ART/'rendered'; TMP=ART/'build_tmp'
for d in (OUT,ART,DOCS,RENDER,TMP): d.mkdir(parents=True,exist_ok=True)
FONT='NanumGothic'; MATH='DejaVu Math TeX Gyre'
NAVY='183A5A'; TEAL='176A70'; INK='223748'; GREY='5D6C7A'; PALE='EDF3F7'
STAMP=datetime.now(timezone.utc).isoformat()
TITLE='R-Delay: 감염병 확산 수리모델 기반 해상·육상 지연 예측·대응 웹서비스'
SUMMARY='해상·육상 지연을 SIS 미분방정식으로 연결하고 로그-OU·입자필터로 변동과 관측을 반영하는 웹서비스를 제안한다. 선석 대기·화물 준비·반출을 구분해 원래 마감 위험과 실행 가능한 배차안을 비교한다. 실자료가 없는 부분은 합성 시나리오로 표시한다.'
MOTIVE='감염병 확산의 전파·회복 구조를 해상·육상 운송의 일정 문제에 적용하고자 지원한다. 공개 운영자료로 문제를 구체화하고, 수리모델·관측 갱신·대기 계산을 웹 의사결정으로 연결한다. 대회에서 모델과 사용자 흐름을 함께 검증하고 데이터 확보·사업화 가설을 구체화하고자 한다.'
KEYS=['SIS 수리모델','로그-OU·입자필터','해상·육상 융합','시간가변 대기모형','배차 의사결정']
EQUATIONS={
'SIS':r'\frac{dI_i}{dt}=\beta_i(t)\left(1-\frac{I_i}{N_i}\right)\sum_j C_{ij}(t)I_j-\gamma_i(t)I_i',
'C':r'C(t)=\mathrm{Id}+\kappa\tau_0 M(t),\qquad S_i=N_i-I_i,\quad x_i=I_i/N_i',
'R':r'R_d(t)=\rho\!\left[\operatorname{diag}\!\left(\beta_i(t)(1-x_i(t))\right)C(t)\Gamma(t)^{-1}\right]',
'Ratio':r'\frac{dx_i}{dt}=\beta_i(1-x_i)\sum_j C_{ij}\frac{N_j}{N_i}x_j-\gamma_i x_i',
'OU':r'dz_i=\vartheta_i(\bar z_i-z_i)\,dt+\sigma_i\,dW_i,\qquad \beta_i(t)=\bar\beta_i(t)e^{z_i(t)}',
'OUstep':r'z_{i,k+1}=\bar z_i+(z_{i,k}-\bar z_i)e^{-\vartheta_i\Delta t}+\sigma_i\sqrt{\frac{1-e^{-2\vartheta_i\Delta t}}{2\vartheta_i}}\,\xi_{i,k}',
'PF':r'\widetilde w_k^{(r)}=w_{k-1}^{(r)}p(Y_k\mid X_k^{(r)}),\qquad w_k^{(r)}=\frac{\widetilde w_k^{(r)}}{\sum_\ell\widetilde w_k^{(\ell)}}',
'ESS':r'\mathrm{ESS}_k=\left[\sum_r\left(w_k^{(r)}\right)^2\right]^{-1}',
'Incident':r'm_{ik}=q_{ik}\int_{t_{k-1}}^{t_k}\beta_i(t)(1-x_i(t))\sum_j C_{ij}(t)I_j(t)\,dt',
'NB':r'Y_{ik}\sim\mathrm{NB}(m_{ik},k_i),\qquad \operatorname{Var}(Y_{ik})=m_{ik}+m_{ik}^2/k_i',
'BB':r'\pi\sim\mathrm{Beta}(\phi p,\phi(1-p)),\qquad Y\mid\pi\sim\mathrm{Binomial}(n,\pi)',
'Queue':r'\dot\pi_n=a(t)\pi_{n-1}+\mu(t)\min(n+1,m)\pi_{n+1}-[a(t)+\mu(t)\min(n,m)]\pi_n',
'Mu':r'\mu(t)=\mu_0(t)\exp[-\eta_B x_P(t)]\,o_B(t)',
'QueueTail':r'\Pr(W>w\mid N=n)=e^{-m\mu w}\sum_{k=0}^{n-m}\frac{(m\mu w)^k}{k!},\qquad n\geq m',
'Ready':r'R_q^{(s)}=\max\{G_{\mathrm{ship}}^{(s)}+\tau_{\mathrm{transfer},q}^{(s)},D_q^{(s)}\}',
'Road':r's_q(a)=\max\{R_q,b_q(a)\},\qquad Q_L(G_q(a))-Q_L(s_q(a))=1',
'Work':r'v_L(t)=\frac{1-\eta_L x_L(t)}{h_0},\qquad Q_L(t)=\int_0^t v_L(u)\,du',
'Arrival':r'A_q(a)=G_q(a)+u_q(a),\qquad W_{\mathrm{truck},q}(a)=\max\{R_q-b_q(a),0\}',
'Risk':r'\widehat p_q(a)=\sum_r w_\tau^{(r)}\frac{1}{S_r}\sum_{s=1}^{S_r}\mathbf{1}\{A_q^{(r,s)}(a)>c_q\}',
'Cost':r'J_q(a)=k_a+c_w E[W_{\mathrm{truck},q}(a)]+L_q\widehat p_q(a)',
'Choose':r'a^*\in\operatorname*{arg\,min}_{a\in\mathcal A_{\mathrm{feasible}},\ \widehat p_q(a)\leq\epsilon}J_q(a)',
'Estimate':r'\widehat\gamma_L=\frac{\int_{t_0}^{t_1}\beta_L(1-y_L)(y_L+c_{LP}y_P)\,dt-[y_L(t_1)-y_L(t_0)]}{\int_{t_0}^{t_1}y_L(t)\,dt}',
'Positive':r'\bar\beta_i(t)=b_i\exp[f_\theta(Z_{\leq t})],\qquad \gamma_i(t)=g_i\exp[h_\psi(Z_{\leq t})]',
'Ident':r'x_P(t)=0.5\Longrightarrow\gamma_P=\beta_P/2;\qquad (\kappa,M)\mapsto(a\kappa,M/a)',
'Brier':r'\mathrm{BS}=\frac{1}{K}\sum_{k=1}^{K}(\widehat p_k-y_k)^2',
'Buffer':r'b^*=\inf\{b\geq0:\Pr(T>t_{\mathrm{plan}}+b)\leq\epsilon\}',
}


def run(args,**kw):
    p=subprocess.run(args,cwd=ROOT,text=True,capture_output=True,**kw)
    if p.returncode: raise RuntimeError('Command failed: '+repr(args)+'\n'+p.stdout+'\n'+p.stderr)
    return p.stdout.strip()


def init_doc(label,compact=False):
    d=Document(); s=d.sections[0]
    s.page_width=Mm(210); s.page_height=Mm(297)
    s.left_margin=s.right_margin=Mm(17)
    s.top_margin=Mm(15);s.bottom_margin=Mm(15)
    s.header_distance=Mm(6);s.footer_distance=Mm(7)
    normal=d.styles['Normal']; normal.font.name=FONT; normal.font.size=Pt(9.2 if compact else 10.5)
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),FONT)
    normal.paragraph_format.line_spacing=1.17 if compact else 1.3
    normal.paragraph_format.space_after=Pt(4 if compact else 7)
    for name,sz in [('Title',23),('Heading 1',15),('Heading 2',11.5)]:
        st=d.styles[name];st.font.name=FONT;st.font.size=Pt(sz);st.font.bold=True;st.font.color.rgb=RGBColor.from_string(NAVY)
        st._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),FONT)
        st.paragraph_format.space_before=Pt(8);st.paragraph_format.space_after=Pt(7)
    d.core_properties.title=label;d.core_properties.author='R-Delay';d.core_properties.subject='v0.1 — 1차 제출·설계 문서 / 서비스 구현 완료본 아님'
    hp=s.header.paragraphs[0];hp.text='R-DELAY  /  '+label+'  /  v0.1'
    hp.runs[0].font.size=Pt(7);hp.runs[0].font.color.rgb=RGBColor.from_string(GREY)
    fp=s.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    rr=fp.add_run('1차 제출·설계 문서  |  ');rr.font.size=Pt(7)
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');fp._p.append(f)
    return d


def para(d,text,size=None,bold=False,style=None):
    p=d.add_paragraph(style=style); r=p.add_run(text);r.bold=bold
    if size:r.font.size=Pt(size)
    return p


def heading(d,text): return para(d,text,style='Heading 1')
def sub(d,text): return para(d,text,style='Heading 2')
def page(d,title,first=False):
    if not first:d.add_page_break()
    heading(d,title)


def shade(cell,fill):
    p=cell._tc.get_or_add_tcPr();s=OxmlElement('w:shd');s.set(qn('w:fill'),fill);p.append(s)


def table(d,rows,widths=None,head=True,size=8.5,heights=None):
    t=d.add_table(rows=len(rows),cols=len(rows[0]));t.style='Table Grid';t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths is None:widths=[176/len(rows[0])]*len(rows[0])
    for col,w in zip(t.columns,widths):col.width=Mm(w)
    for i,row in enumerate(rows):
        trpr=t.rows[i]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if head and i==0:trpr.append(OxmlElement('w:tblHeader'))
        if heights and heights[i]:t.rows[i].height=Mm(heights[i]);t.rows[i].height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
        for j,txt in enumerate(row):
            c=t.cell(i,j);c.width=Mm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            c.text=str(txt)
            tcpr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                x=OxmlElement('w:'+edge);x.set(qn('w:w'),'75');x.set(qn('w:type'),'dxa');mar.append(x)
            tcpr.append(mar)
            if (head and i==0) or (not head and j==0):shade(c,PALE)
            for p in c.paragraphs:
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(2);p.paragraph_format.line_spacing=1.14
                for r in p.runs:r.font.size=Pt(size);r.bold=(head and i==0) or (not head and j==0)
    d.add_paragraph().paragraph_format.space_after=Pt(0)
    return t


EQCACHE={}
def equation(d,key,number=None):
    if key not in EQCACHE:
        md=TMP/(key+'.md');target=TMP/(key+'.docx')
        md.write_text('$$\n'+EQUATIONS[key]+'\n$$\n',encoding='utf-8')
        run(['pandoc',str(md),'-o',str(target)])
        source=Document(target)
        matches=source._element.xpath('.//m:oMathPara')
        if not matches: raise RuntimeError('No native OMML formula for '+key)
        EQCACHE[key]=copy.deepcopy(matches[0])
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(6)
    p._p.append(copy.deepcopy(EQCACHE[key]))
    if number:
        r=p.add_run('   ('+str(number)+')');r.font.size=Pt(8)
    return p


def flow(d,labels):
    rows=[[' → '.join(labels)]]
    t=table(d,rows,[176],head=False,size=9)
    shade(t.cell(0,0),'E8F2F1')
    return t


def application():
    d=init_doc('참가신청서',True)
    para(d,'『2026 동남권 멀티도메인 모빌리티 비즈니스 해커톤』',12,True).alignment=WD_ALIGN_PARAGRAPH.CENTER
    para(d,'참가 신청서 [개인 부문 / 예선]',14,True).alignment=WD_ALIGN_PARAGRAPH.CENTER
    table(d,[['1. 신청 정보','접수번호  '],['팀명','R-Delay'],['제안 작품명',TITLE],['핵심키워드',', '.join(KEYS)],['관련 모빌리티 도메인','융합 대상: ☑ 육상   ☑ 해상   □ 항공\n주 도메인: 해상 ◎'],['제안작품 요약',SUMMARY],['지원동기',MOTIVE]],[31,145],False,9.2,[7,8,13,12,14,29,30])
    para(d,'2. 팀 구성원 정보',10,True)
    table(d,[['구분','성명','출생연도','소속대학/전공학과','학년','학적 상태','휴대전화/이메일'],['팀장','','','','','□재학 □휴학',''],['팀원1','','','','','□재학 □휴학',''],['팀원2','','','','','□재학 □휴학',''],['팀원3','','','','','□재학 □휴학','']],[15,19,19,43,12,28,40],True,8,[9,13,13,13,13])
    para(d,'위와 같이 「2026 동남권 멀티도메인 모빌리티 비즈니스 해커톤」에 참가를 신청합니다. 신청서 및 제출자료의 기재사항에 사실과 다른 내용이 있거나, 귀 대회가 제시한 참가 자격 등의 요건을 충족하지 못할 경우 어떠한 불이익도 감수하겠습니다.',9)
    para(d,'2026년        월        일',9).alignment=WD_ALIGN_PARAGRAPH.CENTER
    para(d,'팀명  R-Delay                 팀장                         (서명)',9).alignment=WD_ALIGN_PARAGRAPH.RIGHT
    para(d,'동남권지역혁신융복합단지추진단 귀하',10,True).alignment=WD_ALIGN_PARAGRAPH.RIGHT
    para(d,'※ 제출된 서류는 반환하지 않으며, 대학원생은 개인부문에 신청이 불가함(기업부문 참여 가능).\n※ 본선기간(26.11.8~11.10) 해외여행(일본)이 가능한 유효 여권 및 출입국 허가 요건을 팀원 전원이 충족하여야 함.',7.5)

    page(d,'해커톤 제안 작품 컨셉 기획안')
    para(d,'1. 제안 작품 요약',11,True)
    table(d,[['팀명','R-Delay'],['제안 작품명',TITLE],['핵심키워드',', '.join(KEYS)],['관련 모빌리티 도메인','융합 대상: ☑ 육상   ☑ 해상   □ 항공 / 주 도메인: 해상 ◎']],[31,145],False,9.2)
    table(d,[['구분','요약'],['기획 배경','부산신항–양산권 수입컨테이너 운송을 적용 후보로 삼는다. BPA의 차량 예약 운영 사례[W1]와 공개된 연결 시뮬레이션 요구[W2]를 참고하여 해상 지연과 내륙 마감의 연결 문제를 다룬다. 실제 고객·화물 연결자료는 미확보다.'],['제안 작품 핵심 컨셉','SIS 미분방정식으로 거점의 지연 전파·회복을 계산하고 로그-OU로 계수 변동, 입자필터로 관측 갱신을 구현한다. 시간가변 선석 대기모형을 화물 준비·반출·도로 이동과 연결해 마감 위험과 대응 비용을 비교한다. 원래 마감과 차량·예약·변경 시한을 유지하며, R_d를 단독 안전기준으로 쓰지 않는다. 원안 식과 추가 연결 가정을 구분한다.'],['최종 성과물 개발 전략 및 계획','이번 대회 목표는 수리모델·계수 추정·웹/API의 통합 서비스, 같은 Python 엔진의 로컬 실행환경, 검증 기록과 GitHub 제출 버전이다. SIS·로그-OU·입자필터·선석 대기·배차 비교를 모두 이번 개발 범위로 둔다. 현재 문서는 1차 설계이며 이 기능들의 구현·실행 검증이 완료된 상태는 아니다.'],['비즈니스 모델 및 사업화 방향','내륙운송을 조율하는 중소 포워더 운영담당자를 사용자 가설로, 운영책임자를 구매자 후보로 둔다. 사업장·관리건수 기준 구독을 검토한다. 공개자료로 업무 가설을 정리하고 연결자료의 이용권 확보 후 제한 파일럿으로 입력 부담·반복 사용·구매자 편익을 확인한다. 가격·고객·매출은 미검증이다.']],[35,141],True,9.3,[8,32,58,44,46])

    page(d,'2. 제안 작품 컨셉 상세 기획안')
    sub(d,'가. 기획 배경')
    para(d,'□ 항만 일정 조정의 문제를 개별 화물의 결정으로 구체화')
    para(d,'BPA는 BPT 신선대·감만의 2021년 11~12월 VBS 시범운영에서 차량 평균 대기시간이 28.23분에서 23.93분으로 감소했다고 발표했다[W1]. 특정 운영 사례이며 R-Delay 성과·부산신항 현재 평균으로 일반화하지 않는다. MSC의 환적 연결 시뮬레이션 요구[W2]는 설계 참고근거이지 중소 포워더 구매 의사의 증거는 아니다.')
    para(d,'□ 기존 기능을 인정하고 추가 판단 가치를 검증')
    para(d,'Port-i는 통합 모니터링·이상 탐지를 제공한다고 안내한다[W3]. 기존 기능이 없다고 주장하지 않고, 화물별 원래 마감·변경조건·비용을 반영한 배차 유지/변경 판단을 검증한다. 사용자는 운송을 조율하고, 실제 예약·차량 배정·추가비용은 권한자가 승인한다.')
    sub(d,'나. 제안 작품 핵심 컨셉 / 나-1. 핵심컨셉')
    para(d,'□ SIS 미분방정식을 이번 구현의 핵심 엔진으로 유지 [U]')
    equation(d,'SIS','1');equation(d,'C','2');equation(d,'R','3')
    table(d,[['기호','뜻·단위'],['Iᵢ, Nᵢ, xᵢ','지연·총 작업 수(공통 작업단위), 지연 비율(무차원). 0≤Iᵢ≤Nᵢ, 단기 고정집단 가정.'],['βᵢ, γᵢ','전파율·지연 회복률(h⁻¹). γ는 선석 처리량과 다름.'],['Mᵢⱼ, τ₀, κ, Cᵢⱼ','j→i 연결 영향률(h⁻¹), 기준시간(h), 배율·결합계수(무차원).'],['Γ, ρ, R_d','회복률 대각행렬, 스펙트럼 반지름, 현재 상태의 전파 지표. 개별 마감 실패확률이 아님.']],[33,143],True,8.3)
    para(d,'모델의 지연 경로를 선석 대기·화물 준비·육상 처리로 전달한다. 지연 작업 수와 선박 대기열 수는 별개이며, 연결식은 추가 가정으로 표시한다. 확정된 원래 납기를 뒤로 옮겨 성공률을 높인 것으로 계산하지 않는다.',8.7)

    page(d,'나-2. 멀티도메인 기술 융합 방안')
    table(d,[['영역','입력·기술','운영 결과로의 연결'],['해상 ◎','입항·작업·지연 관측, 선석 도착·서비스율','SIS 상태·선석 대기 → 대상 선박 작업 → 화물 준비'],['육상','차량 계획, 반출 처리·이동·수령 마감','준비·차량 조건 → 반출·도착 → 실행 가능한 배차 비교']],[23,72,81],True,8.7)
    table(d,[['이번 구현 모듈','역할과 완료 확인'],['로그-OU','전파계수의 양의 확률변동을 생성. 무잡음 기준·난수 재현 검사.'],['입자필터','관측 종류에 맞는 가능도로 상태·가중치 갱신. 누락·불가능 관측 경고.'],['시간가변 대기','선석 시스템의 작업 수 분포와 대상 선박의 대기/완료시간 계산. 확률 보존·범위 절단 점검.'],['계수 추정·학습','현재 이용 가능한 설명변수에서 계수를 추정하고 이후 구간 평가. 상수 추정과 학습 모델 구분.']],[34,142],True,8.7)
    sub(d,'나-3. 컨셉 개념도')
    flow(d,['자료·시나리오','계수 추정·SIS/OU','관측 갱신(PF)','대기·도착','대응 비교'])
    para(d,'사건 연결: 선박 도착 → 선석 대기 → 선박 작업 → 화물 준비 완료 → 차량 인수·반출 → 도로 이동 → 수령지 도착. 선석 대기는 화물 준비시각에 한 번만 반영한다. 웹은 입력·표시, Python API는 동일 계산 엔진을 호출한다.')
    sub(d,'□ 자료 확보·불확실성·실행 실패 대응')
    para(d,'공개 기사·통계는 문제와 업무 설계의 근거이며 개별 학습자료를 대체하지 않는다. 신규 지연 사건 수와 현재 지연 수, 발생시각과 이용 가능시각을 분리한다. 미래 관측·사후 평활화는 실시간 경보에 넣지 않는다. 실측·추정·합성을 표시하고 실제 연결은 ID 대응과 이용권을 확인한다.')
    para(d,'관측·추정 실패 시 이전 추천을 지우고 경고한다. 예약 미확인·차량 중복·변경 시한 경과는 추천에서 제외한다. 표본·운영로그 부족 시 합성자료로 모델 기능을 검증하되 현장 정확도·절감 실적이라고 쓰지 않는다. 로컬 실행은 인터넷 없이 새 계산까지 시험한다.')

    page(d,'다. 최종성과물 개발 전략 및 계획')
    table(d,[['단계','이번 대회 수행 범위·완료 기준'],['사전 준비','공개 근거·관측/단위·입출력 명세, 데이터 이용권·사전개발 허용 범위 확인'],['본선 1일차','SIS·OU·필터·대기·계수 추정 통합, API에서 동일 결과 반환'],['본선 2일차','웹에 관측 전후 상태·위험·후보 사유 표시, 온라인/로컬 실행 점검'],['본선 3일차','해석해·시간간격·확률보존·제약 검사, 시간순 평가·민감도 및 발표 정리']],[30,146],True,8.7)
    para(d,'4인 역할: 수리모델·통합 / 관측·자료·검증 / 웹·시각화 / 공개 업무근거·사업모델·발표. 본선 일정은 첨부 공고 기준이며 변경 공지를 재확인한다[A1]. 계획을 완료 실적으로 표시하지 않는다.')
    para(d,'검증은 해석해·교차 적분·고정 난수·상태 절단, 같은 정보의 단순 기준 비교, 추정에 쓰지 않은 이후 구간으로 구성한다. 위험 개선·비용만 개선·대응 불가·실행 제약 상황을 포함한다. 코드·문서·모델 버전을 일치시키고 실제 실행기록을 보존한다. 지금 단계의 모델 수치시험·웹 배포는 미실행이다.')
    sub(d,'라. 비즈니스 모델 및 사업화 방향')
    para(d,'초기 사용자 가설은 내륙운송 위탁관리 담당자이며 실제 배차권자는 운송사일 수 있다. 자료 조회자·변경 승인자·비용 부담자·구매자를 분리해 검증한다. 구독은 사업장/관리건수 기준으로 검토하되 정해진 매출·가격은 제시하지 않는다.')
    para(d,'공개근거·합성 시제품 → 자료 이용권·연결 확인 → 제한 파일럿 → 입력 부담·반복 사용·구매자 편익 확인 → 유료 전환 순으로 추진한다. 인터뷰를 예선 필수조건으로 두지 않는다. 포워더·운송사·수령지의 정보와 승인 협력 구조를 설계하고 지역 지명만으로 전동·항공 특화기술까지 활용했다고 주장하지 않는다.')
    para(d,'기대효과는 불필요한 대기·변경과 납기 판단 부담의 감소다. 실제 비용·배출저감, 지불의사·협약·고객은 미검증으로 둔다. 실행 가능한 추천이 없어도 잔여 위험을 설명하는 것이 제품 가치다.')
    sub(d,'출처')
    para(d,'[A1] 첨부 모집공고, pp.2–5. [A2] 첨부 참가양식, p.3. [U] 이번 사용자 제공 설계 기준.\n[W1] BPA, VBS 2차 시범운영 결과, 2022.1.27, busanpa.com (dataSid=27214).\n[W2] BPA, MSC와 디지털 전환 협력 강화, 2026.4.13 (dataSid=36141).\n[W3] BPA, Port-i 이용 확대, 2026.3.3 (dataSid=34863). 전체 URL: 종합설계서 출처.',7.5)
    return d


def report():
    d=init_doc('종합설계서')
    para(d,'R-DELAY',12,True)
    para(d,'종합설계서',style='Title')
    para(d,'v0.1 / 1차 제출·설계 문서',14,True)
    para(d,'감염병 확산 수리모델 기반\n해상·육상 지연 예측·대응 웹서비스',16)
    para(d,'수리모델을 화면 뒤의 미정 기술로 남기지 않고,\n관측·대기·화물·대응 판단으로 연결하는 구현 기준.',12)
    table(d,[['이번 문서의 목적','공식 신청서와 일치하는 수식·입출력·검증·사업화 명세'],['이번에 만드는 것','실제 Word/PDF, 문서 생성 원본, 문서 환경·페이지 검사 기록'],['이번에 구현하지 않는 것','서비스·수리모델 계산 엔진 전체. 이들은 이번 대회 구현 목표로 유지'],['자료 상태','실제 고객·연결로그·예측 성능·절감 효과 미검증'],['개인정보','기입하지 않음. 최초 초안의 성명·학력·연구 실적을 복사하지 않음']],[38,138],False,10,[19,23,25,21,23])
    para(d,'작성 기준 2026-10-21. 수치 실적이나 모델 테스트 통과 수를 기재하지 않는다. 이 PDF는 공식 컨셉기획안 4쪽 제한을 우회하는 추가 제출 본문이 아니라 설계·개발·발표 준비자료다.',9)

    page(d,'01  읽는 순서와 상태 표시')
    table(d,[['구간','내용'],['02–05','출처·문제·사용자·범위·사건 정의'],['06–14','SIS·R_d·OU·관측·필터·선석 대기·화물·의사결정·학습'],['15–17','데이터/API·웹/GitHub·검증 및 개발 순서'],['18–20','사업모델·제출 점검·출처와 문서 검사']],[30,146],True,10)
    sub(d,'구현 계획과 완료 사실은 별개다')
    para(d,'설계 작성: 식·모듈·역할을 명세한 상태. 구현: 코드가 작성된 상태. 실행 확인: 해당 코드를 실제 환경에서 수행한 상태. 검증: 명시한 시험의 기대값과 결과를 비교한 상태. 현장 효과: 실제 자료·사용에서 별도로 검토해야 하는 상태다.')
    table(d,[['모듈','이번 대회 범위','이 문서 단계'],['SIS·C·R_d','포함','사용자 식을 보존한 설계'],['로그-OU·입자필터','포함','구현 목표 / 미구현'],['시간가변 선석 대기','포함','구현 목표 / 미구현'],['계수 추정·학습','포함','구현 목표 / 미실행'],['배차·웹/API·로컬 실행','포함','구현 목표 / 미배포']],[55,45,76],True,9.5)
    para(d,'사용자 승인 없이 위 핵심 모듈을 대회 이후 연구로 이동하지 않는다. 개발 순서를 나누는 것과 대회 구현 범위에서 제외하는 것은 다르다. 문서 생성 성공도 수리모델 정확성 검증을 뜻하지 않는다.')

    page(d,'02  자료를 어디에서 확인했는가')
    para(d,'저장소 최초 조회에는 공고·참가양식 PDF가 없었다. 따라서 저장소에서 이 파일들을 읽었다고 표현하지 않는다. 동일 대화에 첨부된 7쪽 공고와 3쪽 참가양식의 본문·표·페이지 이미지를 사용했다. 사용자 핵심식은 이번 메시지[U]를 기준으로 삼는다.')
    table(d,[['근거','확인 범위·용도'],['A1 / 첨부 공고','pp.2–5: 두 도메인+AI·데이터, 자격·일정·예선/본선·제출서류. p.7: 권리·추가공지.'],['A2 / 첨부 양식','pp.1–2: 신청·요약표 이미지. p.3: 상세 항목, 최대4쪽, 출처·도식 허용.'],['U / 현재 사용자 설계','SIS·C·R_d 식, 이번 구현범위, 개인정보 공란, 추가 가정·미구현 상태 구분.'],['최초10쪽 초안','이번 작업의 근거로 원문을 저장소에서 읽었다고 주장하지 않음. 개인실적을 인용하지 않음.']],[42,134],True,10)
    sub(d,'공식 형식과 편집 재량')
    para(d,'신청서: 요약·지원동기 200자 이내, 핵심키워드 5개 이내. 컨셉기획안: 요약 포함 최대4쪽. 가·나 필수이며 다·라도 제공된 공식 항목으로 작성한다. 그림·표·도식을 사용할 수 있고 출처를 기재한다[A2]. 본 문서는 공백 포함 글자 수를 검사한다.')
    para(d,'신청서의 표 순서와 필수 칸을 보존한 편집본이다. 공식 HWP의 픽셀 단위 복제라고 주장하지 않는다. 첨부 양식에 없는 동의·서약 원문을 임의로 다시 쓰거나 증명서를 생성하지 않는다.')

    page(d,'03  공개 근거와 차별화 가설')
    sub(d,'[W1] 차량 도착·대기는 측정 가능한 운영 문제')
    para(d,'BPA가 발표한 BPT 신선대·감만의 과거 VBS 시범운영에서 평균 대기시간은 28.23분에서 23.93분으로 변화했다(2022.1.27). 특정 기간 전후의 발표 결과다. 독립적으로 인과효과를 식별한 실험, R-Delay의 성과, 현재 신항 운송 평균으로 바꾸어 쓰지 않는다.')
    sub(d,'[W2] 실제 이용자의 연결 시뮬레이션 요구')
    para(d,'BPA–MSC 협의 공개자료(2026.4.13)는 환적 연결 시뮬레이션과 터미널 간 운송 통합정보 등의 고도화 의견을 담는다. 이를 기능 설계 참고로 사용한다. 제안된 기능을 이미 구현된 기능으로, 글로벌 선사의 환적 요구를 중소 포워더의 유료 수요로 해석하지 않는다.')
    sub(d,'[W3] 기존 기능을 축소해서 설명하지 않기')
    para(d,'Port-i 공개 안내(2026.3.3)에는 9개 터미널 환적상황 통합과 선박·화물 모니터링/이상 탐지가 있다. 기존 예약·모니터링 체계를 인정하고, 화물별 마감·변경조건·비용을 묶은 판단의 추가 가치를 비교 실험으로 확인한다.')
    table(d,[['자료로 뒷받침되는 범위','아직 뒷받침되지 않는 범위'],['운영 문제와 공개 기능 요구','R-Delay의 고객 구매 의사·협약'],['기존 시스템의 공개된 기능','경쟁사의 모든 기능 조사·독점적 우위'],['문제와 설계 가설','개별 화물 확률분포·연동권·현장 절감']],[88,88],True,9.5)
    para(d,'서로 다른 대기시간·상차시간·선박 지연을 합치지 않는다. 기사 평균 몇 개를 실제 학습자료로 변환하지 않는다. 인터뷰가 어려운 현재 조건에서도 근거의 범위를 지킨다.')

    page(d,'04  사용자·업무·이번 구현 범위')
    table(d,[['역할','설정할 가설','확인할 권한'],['사용자','중소 포워더 운송운영 담당자','준비·일정을 함께 보고 변경 요청을 할 수 있는가'],['실행자','운송사·예약 담당자','차량·기사·터미널 예약을 확정하는가'],['승인자·비용 부담자','계약에 따라 분리 가능','대기·취소·재배차·마감 실패비용 귀속'],['구매자','운영책임자 후보','예산권, 반복 편익, 입력·지원 비용']],[33,58,85],True,9.5)
    para(d,'초기 사례: 부산신항에서 양산권 내륙 수령지로 가는 수입컨테이너 한 건. 실제 구간 자료를 확보한 것으로 표현하지 않는다. 대상은 기본 운송계획이 있지만 변경 검토 시간이 남은 상황이다.')
    sub(d,'입력과 고정 조건')
    para(d,'동일 화물·출발/수령지·원래 마감을 고정한다. 바꾸는 것은 확인 가능한 차량 터미널 도착 계획이다. 예약이 필요한 시간창과 차량의 가용시간은 따로 확인한다. 승인에 걸리는 시간을 제외하고 단순히 변경 마감까지 남은 시간만 표시하지 않는다.')
    sub(d,'이번 단계에서 확장하지 않는 부가 기능')
    para(d,'로그인·결제·복잡한 지도·실제 자동 배차는 첫 통합의 필수조건이 아니다. 수리모델과 관측 갱신·대기·대응 계산은 이번 범위로 유지한다. 항공·지역 확대는 이름만 추가하지 않고 실제 연결과 자료에 따라 판단한다.')

    page(d,'05  사건 정의와 중복 계산 방지')
    flow(d,['선박 도착','선석 대기','선박 작업','화물 준비'])
    flow(d,['화물·차량 모두 준비','인수·반출','도로 이동','수령지 도착'])
    table(d,[['사건','담당 모델·입력','중복 방지 규칙'],['선박 도착','일정/관측 또는 합성 도착 시나리오','입항과 준비시각 구분'],['선석 대기·작업','시간가변 선석 대기열','대상 선박 한 번만 삽입'],['화물 준비','작업 완료+이송, 문서/인도 조건','준비에 포함한 선석 대기 재가산 금지'],['차량 대기·반출','준비와 차량 도착, 육상 처리 진척','별도 트럭 대기열은 기본형에 중복 추가하지 않음'],['도로·수령','출발시각에 따른 이동과 고정 마감','도착/접수/하역 완료 중 평가 사건 명시']],[35,67,74],True,9.5)
    para(d,'SIS 상태의 지연 작업 수와 선석 시스템의 선박 수는 다른 변수다. 모델이 연결된다는 이유만으로 동일 값으로 놓지 않는다. 같은 지연을 준비·서비스율·최종 시간에서 여러 번 세지 않는지 사건 정의서로 검사한다.')
    para(d,'고정한 작업집단은 단기 모형 가정이다. 실제 작업이 유입·이탈하면 정상 복귀와 관측대상 탈락을 구분해야 한다. 이 가정을 검증하지 않고 현실의 개방 물류망을 그대로 표현했다고 주장하지 않는다.')

    page(d,'06  SIS 미분방정식과 결합행렬 [U]')
    equation(d,'SIS',1);equation(d,'C',2)
    table(d,[['기호','의미','단위·조건'],['i, j; t','영향 받는 거점, 원천 거점; 시간','거점 ID; h'],['Iᵢ, Sᵢ, Nᵢ','지연·정상·총 작업 수','공통 의미 작업 수; I+S=N'],['βᵢ, γᵢ','지연 생성·상태 해소 계수','h⁻¹; β≥0, γ>0'],['Mᵢⱼ, τ₀','j→i 영향률, 기준시간','h⁻¹, h'],['κ, Cᵢⱼ','결합 배율·계수','무차원; 비음수'],['xᵢ','Iᵢ/Nᵢ 비율','무차원; 0≤x≤1']],[31,77,68],True,9.5)
    sub(d,'작업집단 크기가 다를 때')
    equation(d,'Ratio',3)
    para(d,'비율 방정식에서는 Nⱼ/Nᵢ가 사라지지 않는다. 첫 구현은 같은 크기의 비교 가능한 작업집단으로 제한할 수 있지만, 다른 크기로 확장할 때 원식을 보존해야 한다. C는 지연 영향의 결합이며 화물 이동량 보존 생성행렬이 아니다.')
    para(d,'Iᵢ=0이면 우변이 비음수, Iᵢ=Nᵢ이면 −γᵢNᵢ이므로 경계는 안쪽을 향한다. 수치해가 범위를 벗어나면 조용히 잘라내지 않고 적분·입력 조건을 점검한다. 초기조건·계수·단위는 결과와 함께 저장한다.')

    page(d,'07  지연 재생산수의 해석 경계 [U]')
    equation(d,'R',4)
    para(d,'Γ=diag(γᵢ)이며 ρ는 스펙트럼 반지름이다. K의 i,j 성분은 새 지연 생성행렬의 i,j 성분을 원천 거점 j의 회복률로 나눈다. F·Γ는 h⁻¹, K·R_d는 무차원이다.')
    table(d,[['구분','표시·해석'],['현재 상태의 R_d','현재 x·계수를 고정한 지표. 전체 미래 안정성 보장 아님.'],['무지연 기준','x=0에서 별도 계산. 현재 지표와 동일시하지 않음.'],['입자별 출력','각 입자에서 계산한 R_d의 가중 요약. 평균 상태의 R와 혼동 금지.'],['화물별 위험','별도 사건·처리 모형의 A>c 확률. R_d를 확률로 대입하지 않음.']],[48,128],True,10)
    sub(d,'단일 임계값으로 경보를 정하지 않는 이유')
    para(d,'한 방향의 연결로 구성된 삼각행렬에서는 비대각 연결항을 바꾸어도 고유값이 같을 수 있다. 그러나 유한 시간의 하류 상태는 변할 수 있다. 따라서 지연 추이·전파 지표·화물 마감 위험을 함께 제공한다. 이는 대수적 구조에 대한 설명이며 실제 운영 성능 수치가 아니다.')
    para(d,'R_d>1만으로 긴급 배차를 추천하거나 R_d≤1을 납기 안전으로 표시하지 않는다. 고유벡터 민감도를 개발한다면 고유값의 단순성·미분가능성 등 적용조건과 인과효과가 아니라는 한계를 기록한다.')

    page(d,'08  로그-OU 확률변동 — 구체화할 설계')
    equation(d,'OU',5);equation(d,'OUstep',6)
    table(d,[['기호','설명','단위·조건'],['zᵢ, z̄ᵢ','전파계수의 무차원 로그 변동·장기 평균','무차원'],['β̄ᵢ, βᵢ','기준 및 변동 전파율','h⁻¹; 양수'],['ϑᵢ','평균회귀 속도','h⁻¹; >0'],['σᵢ','확률변동 크기','h⁻¹ᐟ²; ≥0'],['Wᵢ, ξᵢₖ','브라운 운동·정규 혁신','Var(dW)=dt; ξ~N(0,1)'],['Δt','한 전이 시간간격','h']],[34,78,64],True,9.5)
    para(d,'위 식은 사용자 로그-OU 목표를 구체화한 추가 구현 명세다. 독립 잡음을 초기 기본 가정으로 기록하고, 상관을 넣을 때는 유효한 공분산 구조와 출처를 남긴다. 지수변환된 기준율은 σ>0에서 평균/중앙값이 다르므로 ‘평균 β’라고 자동 호칭하지 않는다.')
    para(d,'OU 끝점 전이는 닫힌 형태로 계산할 수 있지만, 그 사이 SIS 적분은 수치근사다. 시간간격 축소와 무잡음 기준을 별도로 검사한다. 로그정규 랜덤워크를 로그-OU라고 부르지 않는다. OU 잡음만으로 태풍·검역·운영 중단의 발생을 예측했다고 주장하지 않는다.')
    para(d,'난수 시드와 스트림(상태 전이/재표본추출/미래 시나리오)을 분리해 보존한다. 극단적 β에 수치 문제가 생기면 기록 없이 계수를 잘라내지 말고 실패·한계 상태를 반환한다.')

    page(d,'09  관측 정의 — 사건 수와 현재 상태 수')
    equation(d,'Incident',7);equation(d,'NB',8)
    para(d,'신규 지연 사건 수 Yᵢₖ에는 구간 내 정상→지연 생성항의 적분을 평균으로 연결한다. qᵢₖ는 보고/관측 비율(0~1), kᵢ는 음이항 크기(>0)다. 반복 지연 사건은 집단 크기를 넘을 수 있으며 미시적 전환 구조가 없는 평균장 SIS에서는 이 연결 자체가 관측 가정이다. m=0은 0건 퇴화분포로 처리한다.')
    equation(d,'BB',9)
    para(d,'총 n건 중 현재 지연 수에는 상한을 지키는 베타-이항을 추가 설계로 사용한다. p는 입자의 지연 비율, φ>0은 집중도, π는 잠재 확률이다. p=0/1은 경계 퇴화분포로 별도 처리한다. 음이항을 다른 분포로 몰래 교체하지 않고 관측 종류를 API에 명시한다.')
    table(d,[['관측','필드·처리'],['구간 신규 사건 수','구간 시작/종료, 건수, 보고비율, 중복 집계 여부'],['현재 지연 표본 수','event_time, 현재지연 Y, 표본총수 n'],['누락/지연 보고','available_at 별도 저장. 누락은 0건이 아님.'],['같은 자료의 중복 사용','발생/상태 관측의 조건부 독립을 검토. 독립으로 보고 우도 두 번 곱하지 않음.']],[45,131],True,9.5)
    para(d,'이전 대화의 지정 ±0.5%p 사례는 실제 건수 관측이나 신뢰구간이 아니다. 이번 문서에서는 미실행된 복원 수치를 검증 실적으로 싣지 않는다.')

    page(d,'10  입자필터의 예측·갱신·기록')
    equation(d,'PF',10);equation(d,'ESS',11)
    table(d,[['기호','의미'],['Xₖ⁽ʳ⁾','입자 r의 SIS·로그-OU 상태와 연결된 대기열 상태/분포'],['wₖ⁽ʳ⁾','관측 후 정규화 가중치. 합1.'],['Yₖ','선택한 종류의 현재 이용 가능한 관측'],['P, ESS','입자 수, 가중치 퇴화 진단. 1≤ESS≤P.']],[42,134],True,10)
    para(d,'개발 시작 설정은 입자256개, ESS<P/2에서 체계적 재표본추출이다. 이는 성능을 입증한 최적값이 아니라 조정 가능한 구현 가정이다. 로그가능도에서 안정적으로 정규화하고 모든 입자가 사실상 불가능한 관측이면 정상 갱신으로 위장하지 않는다.')
    flow(d,['OU·SIS 예측','관측 가능도','가중치 갱신','ESS 검사','필요 시 재표본추출'])
    para(d,'재표본추출 시 상태뿐 아니라 같은 입자에 조건부인 대기열 분포/시뮬레이션 이력도 함께 복사한다. 다만 아직 태깅하지 않은 미래 대상 선박을 배경 시스템에 중복 삽입하지 않는다. 관측 갱신 요약은 재표본추출 전 가중 분포 기준으로 저장한다.')
    para(d,'실시간은 available_at≤판단시각인 자료만 사용한다. 사후 평활화·중심 이동평균의 미래 정보는 실시간 화면에 넣지 않는다. 늦게 도착한 과거 관측은 원시시각부터 재처리하거나 재실행 필요 상태로 기록한다.')

    page(d,'11  시간가변 선석 대기열')
    equation(d,'Queue',12);equation(d,'Mu',13)
    table(d,[['기호','의미','단위·가정'],['πₙ(t), n','서비스 중 포함 선박 수가 n일 확률','n=0,1,…; 합1; π₋₁=0'],['a(t)','배경 선박 도착률','선박/h; 비정상 포아송 가정'],['μ₀(t), μ(t)','선석 한 개의 서비스 위험률','h⁻¹; 동질 서버·지수 서비스'],['m','동시에 서비스 가능한 선석 서버 수','양의 정수; 고정 기본형'],['η_B, o_B(t)','SIS 서비스 영향·개방 상태','무차원; η_B≥0, o_B=0 또는1']],[35,77,64],True,9.3)
    para(d,'SIS 상태와 서비스율의 관계 μ=μ₀exp(−η_Bx_P)o_B는 추가 가정이다. γ와 μ는 같은 양이 아니며 선석 서비스율 저하를 다시 독립 고정 지연으로 더하지 않는다. 도착률은 별도 일정/관측으로 정한다. m 자체가 바뀌거나 선박 이질성·우선순위가 있으면 상태/서비스 규칙을 확장해야 한다.')
    para(d,'수치계산은 π₀부터 유한 상한까지 하되 상단 도달/초과의 누적 확률을 별도로 추적한다. 마지막 상태에 확률을 몰아넣은 채 절단 오류를 숨기지 않는다. 범위를 늘렸을 때 π·대기위험이 안정되는지 검사한다. 서비스 중단은 μ=0으로 구간 정의하고 재개·잔여작업 처리 규칙을 명시한다.')

    page(d,'12  대기열 분포에서 대상 선박 시간으로')
    para(d,'전진방정식이 주는 값은 선박 수 분포이지 대기시간 분포 자체가 아니다. 대상 선박 도착 직전의 시스템 상태를 입력받아 선입선출 규칙으로 대상의 서비스 시작·완료시각을 추적한다. 대상 선박이 배경 도착률에 포함되었는지 확인해 한 번만 삽입한다.')
    sub(d,'단순 조건에서 사용할 검산식 — 설계용 도출')
    equation(d,'QueueTail',14)
    para(d,'이 식은 대상 도착 시 n척, 서버 m개 고정, 선입선출, 대상의 대기 중 μ 일정·지수 서비스라는 조건에서 n≥m일 때 적용한다. 대상 서비스가 시작될 때까지 n−m+1번의 완료를 기다린다는 조건부 꼬리다. n<m이면 W=0이다. 현재 시간가변 운영 전체에 무조건 적용하지 않는다.')
    table(d,[['경계 상황','계획한 검사'],['빈 시스템·서버 여유','즉시 서비스 시작'],['μ=0인 중단 구간','작업 완료 없음; 재개 후 조건 반영'],['시간가변 μ','누적 위험 또는 사건 시뮬레이션으로 시간 변화 반영'],['계산기간 이후 미완료','검열·미완료 확률 보존. 마지막 시각에 완료한 것으로 처리하지 않음.']],[45,131],True,10)
    para(d,'선박 한 척의 평균 작업시간을 모든 컨테이너의 실제 반출시각으로 일반화하지 않는다. 아래 준비 연결은 가장 단순한 시연 가정이며, 실제 자료가 생기면 컨테이너별 하역·인도 사건으로 바꿀 수 있게 입력 구조를 분리한다.')

    page(d,'13  화물 준비·육상 처리의 추가 연결식')
    equation(d,'Ready',15);equation(d,'Work',16);equation(d,'Road',17);equation(d,'Arrival',18)
    table(d,[['기호','뜻·단위·조건'],['R_q, G_ship, D_q','화물 준비·대상 선박 작업 완료·문서/인도 조건시각(h)'],['τ_transfer,q','작업 완료 후 이송·정리시간(h). 아직 실측되지 않음.'],['b_q(a), s_q(a)','차량 터미널 도착·처리 시작시각(h)'],['h₀, η_L','지연 없는 인수반출 시간(h), 영향계수(0≤η_L<1)'],['v_L, Q_L, G_q','처리 진척률(h⁻¹), 누적 정규화 작업량, 반출 완료(h)'],['u_q(a), A_q(a)','반출 후 이동시간(h), 수령지 도착(h)']],[48,128],True,9)
    para(d,'모든 식은 [U]에서 제공한 SIS 식과 구분되는 추가 가정이다. 화물 준비 R에 선석 대기를 이미 포함했으므로 최종 도착 A에 다시 더하지 않는다. 차량의 준비 대기는 다른 자원의 대기이며 비용으로 별도 기록한다. 처리 진척을 쓴 기본형에 별도의 동일 트럭 처리 대기열을 중복 추가하지 않는다.')
    para(d,'도로 이동시간은 바뀐 반출시각과 경로에 조건부로 계산하는 것을 목표로 한다. 자료가 없어 상수로 두면 그 단순화를 표시한다. 한 화물의 배차 선택이 전체 SIS를 바꾸지 않는 비피드백 가정을 사용하고 공동 최적화 실적으로 표현하지 않는다.')

    page(d,'14  마감 위험·비용·실행 가능성')
    equation(d,'Risk',19);equation(d,'Cost',20);equation(d,'Choose',21)
    table(d,[['기호','설명'],['r, s; w_τ⁽ʳ⁾','갱신된 입자·그 입자의 미래 표본, 판단시점 가중치'],['S_r, c_q, ε','입자별 미래 표본 수, 원래 수령지 마감, 허용 실패확률'],['k_a, c_w, L_q','변경비용(원), 준비 대기단가(원/h), 마감 실패비용(원)'],['A_feasible','예약·차량·변경 시한·승인이 확인된 후보 집합']],[48,128],True,9.5)
    para(d,'추천을 만들려고 ε나 c_q를 사후 변경하지 않는다. 정확히 마감 도착을 성공으로 볼지, 접수·하역 완료를 볼지 미리 고정한다. 실행 가능하다고 납기를 보장하지 않으며, 싼 후보가 위험기준을 못 지키면 추천하지 않는다.')
    para(d,'예약 미확인 후보는 확인 필요로 표시한다. 변경 마감에서 검토·승인·접근 시간을 차감한 대응 가능시간을 보여준다. 비용은 귀속 주체와 실제 단가를 확인하기 전 가정이다. 마감 이후까지 미완료인 표본을 원래 마감 성공으로 처리하지 않는다. 비용 꼬리 평가가 안 되면 최적 비용 추천을 확정하지 않는다.')
    equation(d,'Buffer',22)
    para(d,'사전계획의 필요한 버퍼 b*와 이미 확정된 원래 마감을 구분한다. T는 계획 대상 도착/소요시간, t_plan은 사전 계획 기준이다. 계획 이후 마감을 늦춘 결과를 원래 약속 달성률 개선으로 세지 않는다.',9)

    page(d,'15  계수 추정·학습과 식별 가능성')
    equation(d,'Estimate',23)
    para(d,'이는 β_L·c_LP가 알려졌고 N_P=N_L인 축소모형의 적분 기반 상수 회복률 추정 설계다. y는 관측 지연 비율이며 원안의 머신러닝 학습식 자체가 아니다. 실제 관측을 읽고 추정한 실행 결과는 이번에 제시하지 않는다. 분모가 작거나 비양수 계수가 나오면 실패로 보고한다.')
    equation(d,'Positive',24)
    para(d,'이번 학습 구현 목표는 예측 당시 이용 가능한 설명변수 Z에서 양의 β·γ를 내는 작은 계수 모델을 SIS 안에 넣는 것이다. bᵢ·gᵢ는 단위를 가진 기준율, f·h 출력은 무차원이다. 모든 계수와 연결을 동시에 자유롭게 학습하지 않고 고정/추정 목록과 식별 조건을 정한다.')
    equation(d,'Ident',25)
    para(d,'평형 궤적에서는 β와 γ의 비만 식별될 수 있고 κ와 M의 배율도 구별되지 않는다. 비평형·다중 조건의 합성 복원부터 확인하되, 같은 모델로 자료를 생성하고 다시 적합한 결과를 현장 외삽 능력으로 표현하지 않는다.')
    para(d,'학습·보정·평가를 시간순으로 나눈다. 사후 확정 ETA·미래 기상·평활화·중심 이동평균을 당시 이용 가능 입력으로 쓰지 않는다. 추정값은 이후 구간 또는 동일 계수가 유지된다는 다음 에피소드에 적용한다. 자료 생성 정답은 검산에만 쓰고 추정 함수에 전달하지 않는다.')

    page(d,'16  데이터·API 계약')
    table(d,[['자료군','필수 기록'],['출처·시각','source_id, data_type(실측/추정/합성), event_time, available_at'],['관측','observation_type, 값, 구간 또는 표본 수, 누락 상태'],['화물·실행','익명 ID·선박 연결 ID·원래 마감 사건·차량·예약·시한·승인'],['재현성','model_version, parameter_set_id, commit, 시드·스트림·실행환경']],[40,136],True,9.5)
    table(d,[['API 설계','입력·반환'],['POST /api/v1/calibrations','관측/고정조건 → 추정 모델 ID·진단·학습 구간'],['POST /api/v1/runs','시나리오/모델/마감/후보 → 실행 ID·처리 상태'],['GET /api/v1/runs/{id}','SIS·R_d·대기·위험·후보 및 경고'],['POST /api/v1/runs/{id}/observations','새 관측 → 원본을 보존하는 새 실행 ID'],['GET /api/v1/runs/{id}/artifacts','설정·결과·진단 기록 저장']],[83,93],True,9)
    para(d,'외부 시각은 시간대 포함 ISO 형식, 내부는 한 기준시점 이후 h로 변환한다. 분당/시간당 비용을 혼용하지 않는다. 긴 계산은 실행 ID로 진행상태를 조회하고 중복 요청을 구분한다. 응답에는 합성 여부·모델 상태·계산 경고를 필수로 포함한다.')
    para(d,'계산 실패 / 입력 부족 / 확인 필요 / 기준 충족 후보 없음 / 추천 있음은 서로 다른 상태다. 추정·API 오류가 나면 옛 추천을 현재 결과처럼 남기지 않는다. 늦은 과거 관측은 재처리 규칙과 시간을 보존한다. 위 API는 이번 개발 명세이며 현재 배포된 주소가 아니다.')

    page(d,'17  웹·GitHub·로컬 실행 및 검증 순서')
    table(d,[['화면·폴더','담당 역할'],['상황 / web/','화물·마감·자료 상태·후보 입력'],['모델 분석 / engine/','SIS 수식·단위·상태·R_d·확률변동 표시'],['관측 갱신 / api/','입자 분포·ESS·갱신 전후 및 경고'],['운영 판단','선석 대기·준비·도착·비용·후보 제외 이유'],['근거 / docs/·tests/','실행·모델·코드 버전과 실제 검사 기록']],[60,116],True,9.5)
    para(d,'Python 엔진은 한 곳에서 계산하고 웹은 API 응답을 표시한다. React/TypeScript·FastAPI·컨테이너 구성을 개발 후보로 삼되 지금 설치·배포했다고 표현하지 않는다. 온라인과 로컬은 같은 엔진을 사용하고 캐시 화면이 아니라 인터넷 없이 새 계산이 되는지 확인한다.')
    para(d,'GitHub에는 engine/api/web/tests/experiments/docs와 합성자료를 분리한다. 공개 저장소에 서명·증명서·연락처·비밀키를 넣지 않는다. 코드·문서·실행 버전을 고정하고 검토 없이 main을 병합하거나 공개설정을 바꾸지 않는다.')
    table(d,[['순서','완료 관문'],['1','입력/단위·SIS·OU → 관측1회 PF → API 동일 결과'],['2','선석 대기와 대상 선박 추적 → 준비·육상 처리 → 후보 평가'],['3','웹 한 사례의 입력→갱신→결정→근거 저장'],['4','해석해·교차적분·시간간격·난수·확률보존·미완료 표본·제약 검사']],[19,157],True,9.3)
    para(d,'초기 검수 계획: 상태/확률 범위, 상단 절단 질량, 입력 경계, 고정 난수, 관측 누락, 실행 불가, API/저장 일치. 수치오차와 표본오차를 분리한다. 현재 이 모델 검사들은 미실행이며 문서 검사만 별도로 수행한다.')

    page(d,'18  사업모델·평가·지역 연계')
    table(d,[['단계','검증할 조건'],['공개근거·시제품','실제 업무와 다른 가정 식별, 이번 대회 모델/웹 통합'],['자료 협의','이용권·화물 연결·시간 이력·입력 부담'],['제한 파일럿','반복 사용·추가 판단·권고 수용·오류 대응'],['구독 검토','구매자에 남는 편익·지원원가·유료 전환 의사']],[37,139],True,10)
    para(d,'초기 수익모형은 사업장 또는 관리건수 기준 구독으로 집중한다. 공개기사의 대기시간 감소를 자체 절감률로 쓰지 않고, 전체 사회비용과 고객이 실제 부담하는 비용을 구분한다. 가격·고객 수·매출·시장 규모는 확인 전에 임의로 채우지 않는다.')
    para(d,'지역 협력은 항만 준비정보 제공자, 운송사 차량·일정 제공자, 내륙 수령조건 제공자와 승인 주체의 역할로 설명한다. 부산·양산 지리만으로 울산 전동·경남 항공 특화기술을 모두 결합했다고 말하지 않는다. 향후 확장은 실제 산업 수요와 자료를 전제로 한다[A1 p.2].')
    equation(d,'Brier',26)
    para(d,'BS는 예측 실패확률과 실제 실패 라벨의 제곱오차 평균이다. K는 평가 건수, y∈{0,1}. 본선 검증계획에는 확률 보정·경보 선행시간·오경보·소요시간도 포함한다. 같은 사전 정보와 고정 마감/비용의 기준모형으로 비교한다. 아직 실제 평가값은 없다.')
    para(d,'기대효과는 불필요한 대기·변경·조회 부담 감소다. 환경 편익·현장 절감·회복탄력성 향상은 검증 가설이다. 문헌 근거와 합성실험은 구매 의사나 실제 성능을 대신 입증하지 않는다.')

    page(d,'19  제출 전 확인과 문서 완료 범위')
    table(d,[['항목','처리·직접 확인'],['신청서/컨셉','신청서1쪽, 컨셉 요약 포함4쪽. 요약/동기 200자 이하, 키워드5개 이하.'],['개인정보','성명·생년·학교/전공·학년·연락처·학적 선택·서명 공란. 월일 공란.'],['팀·증빙','학부 재학/휴학 4인, 전원 재학/휴학증명, 출입국 요건 직접 확인[A1].'],['별도 서류','개인정보/초상권 동의, 참가서약, 보안서약은 공식 원본에 직접 작성.'],['최신 공지','첨부 공고 일정은 변경 가능. 추가 규칙·사전 개발 허용 범위는 제출 직전 확인.'],['권리·공개','원안·코드 라이선스·자료이용권 확인. 개인정보 공개 금지.']],[39,137],True,9.8)
    para(d,'이 파일은 제출 문안과 상세 설계의 1차본이다. 서비스·수리모델 실행·배포·현장 성과는 완료 사실로 기재하지 않았다. 과거 대화의 복원값·마감위험 수치·테스트 통과 기록은 이 단계의 성과로 인용하지 않았다.')
    para(d,'문서 생성 시 실제 환경·파일 크기·쪽수·OMML 수식·문자 수·공란을 검사하고 모든 페이지를 PNG로 렌더링한다. 자동검사는 육안검수를 대신하지 않는다. 확인된 육안 범위는 artifacts/document_check.md에 별도로 기록한다.')
    para(d,'환경 제한: 이 대화의 로컬 컨테이너/Python 시작 실패를 원격 실행 성공과 혼합하지 않는다. 저장소의 전용 문서 워크플로에서 실행한 내용만 원격 실행 기록으로 남긴다. 저장소에 없던 첨부 PDF를 저장소에서 읽었다고 쓰지 않는다.')

    page(d,'20  출처와 가정의 출처 분리')
    sub(d,'[A1] 사용자 첨부 모집공고 / 7쪽')
    para(d,'「멀티도메인 모빌리티 비즈니스 해커톤.pdf」. pp.2–5의 주제·자격·서식·일정, p.7의 권리/추가공지. 동일 대화의 본문과 페이지 이미지를 확인했다. 저장소에는 원본 PDF가 없었다.',9.5)
    sub(d,'[A2] 사용자 첨부 참가양식 / 3쪽')
    para(d,'「참가양식.pdf」. pp.1–2 표 이미지, p.3 작성방법·상세 항목. 개인정보 칸과 요약표 4개 행, 컨셉 최대4쪽을 기준으로 삼았다.',9.5)
    sub(d,'[U] 이번 사용자 제공 설계 기준')
    para(d,'SIS·C·R_d 식과 이번 대회 범위는 사용자 메시지를 기준으로 보존했다. 최초10쪽 초안 전체를 저장소에서 읽거나 개인 논문 실적을 확인한 것으로 표현하지 않는다.',9.5)
    for code,sid,title,date in [('W1','27214','VBS 2차 시범운영 결과','2022-01-27'),('W2','36141','MSC와 디지털 전환 협력 강화','2026-04-13'),('W3','34863','Port-i 항만이용자 전체 확대','2026-03-03')]:
        sub(d,'['+code+'] 부산항만공사 / '+title+' / '+date)
        para(d,'https://www.busanpa.com/board/view.bpa?boardId=BBS_0000031&dataSid='+sid+'&menuCd=DOM_000000105002001000',7.5)
    para(d,'외부 원문 확인 기준일 2026-10-21. 전체 주장·적용 범위는 docs/source_audit.md. 모델의 새로운 관측·준비·처리식과 API·검증 기준은 설계 제안/수학적 관계이며 외부 실증 결과가 아니다. 최신 대회 홈페이지 접근 실패를 확인 성공으로 기록하지 않았다.',9)
    return d


def convert(docx,pdf_dir):
    profile=TMP/('lo-profile-'+docx.stem)
    output=run(['libreoffice','-env:UserInstallation='+profile.resolve().as_uri(),'--headless','--convert-to','pdf','--outdir',str(pdf_dir),str(docx)],timeout=180)
    dest=pdf_dir/(docx.stem+'.pdf')
    if not dest.exists() or dest.stat().st_size<5000:raise RuntimeError('PDF conversion did not create a valid file: '+output)
    return dest


def inspect_pdf(path,tag):
    pdf=fitz.open(path); info=[]; images=[]
    for i,p in enumerate(pdf):
        text=p.get_text(); bbox=[]
        for b in p.get_text('dict')['blocks']:
            for line in b.get('lines',[]):
                for s in line.get('spans',[]):
                    if s['text'].strip() and (s['bbox'][0]<-1 or s['bbox'][1]<-1 or s['bbox'][2]>p.rect.width+1 or s['bbox'][3]>p.rect.height+1):bbox.append(s['text'])
        image=RENDER/(tag+f'-{i+1:02d}.png')
        p.get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False).save(image);images.append(image)
        info.append({'page':i+1,'text_characters':len(text),'outside_page_spans':bbox,'replacement_characters':text.count('\ufffd'),'has_korean':bool(re.search('[가-힣]',text))})
        (ART/(tag+f'-{i+1:02d}.txt')).write_text(text,encoding='utf-8')
    # Contact sheet only arranges actual rendered page images; not a replacement for individual review.
    thumbs=[]
    font_path='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
    font=ImageFont.truetype(font_path,16)
    for i,file in enumerate(images):
        im=Image.open(file).convert('RGB');im.thumbnail((365,520))
        cell=Image.new('RGB',(385,557),'#e6ebf0');cell.paste(im,((385-im.width)//2,25))
        ImageDraw.Draw(cell).text((12,535),f'{tag} / {i+1}',fill='#203344',font=font);thumbs.append(cell)
    for batch in range(0,len(thumbs),8):
        group=thumbs[batch:batch+8];sheet=Image.new('RGB',(385*4,557*2),'white')
        for j,im in enumerate(group):sheet.paste(im,((j%4)*385,(j//4)*557))
        sheet.save(RENDER/(tag+f'-contact-{batch//8+1}.jpg'),quality=90)
    return {'file':str(path.relative_to(ROOT)),'pages':len(pdf),'page_checks':info,'rendered_pages':[str(x.relative_to(ROOT)) for x in images]}


def main():
    before=run(['git','status','--short','--branch']);branch=run(['git','branch','--show-current']);source_commit=run(['git','rev-parse','HEAD'])
    if branch!='codex/rdelay-bootstrap':raise RuntimeError('Only the user-approved document branch may be used.')
    skills={}
    for item in ('docx','pdfs'):
        p=Path('/home/oai/skills')/item/'SKILL.md'
        skills[item]={'present':p.exists()}
        if p.exists():skills[item]['read_characters']=len(p.read_text())
    env={'timestamp_utc':STAMP,'branch':branch,'source_commit':source_commit,'git_status_before':before,'python':sys.version,'platform':platform.platform(),'git':run(['git','--version']),'pandoc':run(['pandoc','--version']).splitlines()[0],'libreoffice':run(['libreoffice','--version']),'render_engine':'PyMuPDF '+fitz.VersionBind,'font':run(['fc-match',FONT]),'skills':skills,'local_chat_runtime':'container/Python initialization failed before commands; remote workflow used','service_tests':'not run'}
    probe=ART/'write_probe.tmp';probe.write_text('RDelay UTF-8 파일 쓰기 확인',encoding='utf-8');assert probe.read_text()=='RDelay UTF-8 파일 쓰기 확인';probe.unlink();env['file_write_read_delete_probe']='pass'
    (ART/'environment_check.json').write_text(json.dumps(env,ensure_ascii=False,indent=2),encoding='utf-8')
    (ART/'environment_check.md').write_text('# 실행환경 확인\n\n'+ '\n'.join(f'- {k}: {v}' for k,v in env.items())+'\n\n저장소에는 공고·양식 PDF가 없었으며 대화 첨부자료를 사용했다. 로컬 skill 경로는 접근 시도했으나 런타임 초기화 오류로 읽지 못했다. 원격 경로도 위 present/read 기록 외 읽었다고 주장하지 않는다.\n',encoding='utf-8')
    assert len(SUMMARY)<=200 and len(MOTIVE)<=200 and len(KEYS)<=5
    app_path=OUT/'RDelay_참가신청서_v0.1.docx';design_path=OUT/'RDelay_종합설계서_v0.1.docx'
    application().save(app_path);report().save(design_path)
    apppdf=convert(app_path,OUT);designpdf=convert(design_path,OUT)
    app_info=inspect_pdf(apppdf,'application');design_info=inspect_pdf(designpdf,'design')
    d=Document(app_path)
    omissions=[]
    # Team-information table is the second table; never infer blanks from text search alone.
    team=d.tables[1]
    for ri in range(1,5):
        for ci in (1,2,3,4,6):
            if team.cell(ri,ci).text.strip():omissions.append(f'team row {ri} column {ci} not blank')
        if team.cell(ri,5).text!='□재학 □휴학':omissions.append('학적 선택 변경')
    equations=len(d._element.xpath('.//m:oMath'))
    whole='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
    raw_tex=bool(re.search(r'\\(?:frac|beta|gamma|sum|operatorname)',whole))
    checks={'generated_at':STAMP,'source_commit':source_commit,'application':app_info,'design':design_info,'summary_characters_including_spaces':len(SUMMARY),'motivation_characters_including_spaces':len(MOTIVE),'keyword_count':len(KEYS),'native_omml_count':equations,'team_blank_errors':omissions,'raw_latex_in_visible_docx_text':raw_tex,'visual_review':'not yet performed by assistant; actual per-page renders generated','model_tests':'not run','web_deployment':'not performed','application_target_pages':5,'design_target_pages':21}
    # Map requested section starts to pages; spillover is a failed layout test, not a hidden page count change.
    markers=['참가 신청서','해커톤 제안 작품 컨셉 기획안','2. 제안 작품 컨셉 상세 기획안','나-2. 멀티도메인 기술 융합 방안','다. 최종성과물 개발 전략 및 계획']
    with fitz.open(apppdf) as f:
        checks['application_section_pages']={m:[i+1 for i,p in enumerate(f) if m in p.get_text()] for m in markers}
    checks['automatic_pass']=app_info['pages']==5 and design_info['pages']==21 and equations>=3 and not omissions and not raw_tex and all(not p['outside_page_spans'] for report_info in (app_info,design_info) for p in report_info['page_checks'])
    checks['files']=[{'path':str(p.relative_to(ROOT)),'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in sorted(OUT.iterdir()) if p.is_file()]
    (ART/'document_check.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
    md='# 문서 검사 — 실제 실행 기록\n\n'
    md+=f'- 생성 입력 커밋: `{source_commit}`\n- 신청서 PDF: {app_info["pages"]}쪽 (목표5=신청1+컨셉4)\n- 종합설계 PDF: {design_info["pages"]}쪽 (목표21)\n- 요약: {len(SUMMARY)}자 / 지원동기: {len(MOTIVE)}자 (공백 포함)\n- 키워드: {len(KEYS)}개\n- Word 네이티브 OMML 수식: {equations}개\n- 팀정보 공란 오류: {omissions}\n- 본문 LaTeX 소스 노출: {raw_tex}\n- 페이지 전체 렌더링: 완료\n- 자동 범위검사: {checks["automatic_pass"]}\n- 육안검수: 아직 미실행. 자동 검사 통과를 육안검수로 기록하지 않음.\n- 모델 수치시험·웹/API 실행·현장 성능: 미실행\n\n'
    md+='## 파일\n\n'+''.join(f'- `{x["path"]}` — {x["bytes"]:,} bytes\n' for x in checks['files'])
    md+='\n렌더 이미지: artifacts/rendered/. 상세 문자·경계·공란·쪽수: document_check.json. 공식 HWP 원본의 픽셀 단위 복제나 서명·증빙 완료를 의미하지 않는다.\n'
    (ART/'document_check.md').write_text(md,encoding='utf-8')
    (DOCS/'implementation_status.md').write_text('# 구현 상태 — v0.1\n\n## 완료\n- 사용자 설계 기준과 공식 첨부 양식 대조 문안\n- 실제 신청서 Word 및 상세 설계 PDF 생성\n- 문서 변환·쪽수·수식·공란·페이지 경계 자동 검사 및 전체 페이지 렌더링\n\n## 별도 확인\n- 육안검수는 artifacts/document_check.md의 최신 기록 참조\n- 개인정보 직접 기입, 공식 동의·서약서, 재학/휴학 증명, 자격·사전개발 규칙\n\n## 이번 대회 구현 목표 / 아직 미구현·미실행\nSIS·C·R_d 엔진, 로그-OU, 입자필터, 시간가변 선석 대기와 대상선박, 계수 추정/학습, 배차 판단, 웹/API·온라인/로컬 배포. 문서 생성 시험은 이 기능의 수치/성능 검증이 아니다. 기존 대화의 미실행 수치를 실적으로 쓰지 않는다.\n\nmain 자동 병합·다른 연구 저장소 수정·공개 전환·유료 결제·외부 서비스 배포는 하지 않았다.\n',encoding='utf-8')
    # Keep requested reproducibility sources; package no fonts or private source attachments.
    package=OUT/'RDelay_v0.1_문서및생성원본.zip'
    with zipfile.ZipFile(package,'w',zipfile.ZIP_DEFLATED) as z:
        for base in (DOCS,ROOT/'scripts'):
            for p in base.rglob('*'):
                if p.is_file() and '__pycache__' not in str(p):z.write(p,str(p.relative_to(ROOT)))
        z.write(ROOT/'AGENTS.md','AGENTS.md')
        for p in (app_path,apppdf,design_path,designpdf,ART/'environment_check.md',ART/'document_check.md',ART/'document_check.json'):
            z.write(p,str(p.relative_to(ROOT)))
    print(json.dumps(checks,ensure_ascii=False,indent=2))
    if not checks['automatic_pass']:
        print('LAYOUT REVIEW REQUIRED: artifacts preserved; do not mark final submission ready.',file=sys.stderr)

if __name__=='__main__':main()
